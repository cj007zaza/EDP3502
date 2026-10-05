import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_callout_border(cell, border_hex="2563EB", border_size="32"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def set_table_borders(table, border_color="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{border_color}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_run_font(run, font_name="TH Sarabun New", size_pt=14, bold=False, italic=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    # Set complex script font for Thai
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)

def build_chapter8_doc():
    doc = docx.Document()

    # 1 inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

    # Base style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'TH Sarabun New'
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Header Banner Box
    t_banner = doc.add_table(rows=1, cols=1)
    t_banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_banner.autofit = False
    t_banner.columns[0].width = Inches(6.5)
    c_banner = t_banner.cell(0, 0)
    set_cell_background(c_banner, "1E3A8A")
    set_cell_margins(c_banner, top=200, bottom=200, left=240, right=240)

    p_b0 = c_banner.paragraphs[0]
    p_b0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_b0.paragraph_format.space_after = Pt(2)
    r_b0 = p_b0.add_run("รายวิชา EDP3502 การช่วยเหลือและให้คำปรึกษาผู้เรียน")
    set_run_font(r_b0, "TH Sarabun New", 13, bold=False, color_rgb=RGBColor(191, 219, 254))

    p_b1 = c_banner.add_paragraph()
    p_b1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_b1.paragraph_format.space_after = Pt(4)
    r_b1 = p_b1.add_run("สรุปเนื้อหาบทที่ 8 : การให้คำปรึกษาทฤษฎีแบบจิตวิเคราะห์")
    set_run_font(r_b1, "TH Sarabun New", 20, bold=True, color_rgb=RGBColor(255, 255, 255))

    p_b2 = c_banner.add_paragraph()
    p_b2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_b2.paragraph_format.space_after = Pt(0)
    r_b2 = p_b2.add_run("Psychoanalytic Counseling Theory (Sigmund Freud)")
    set_run_font(r_b2, "TH Sarabun New", 14, italic=True, color_rgb=RGBColor(226, 232, 240))

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(8)
    p_sp.paragraph_format.space_after = Pt(4)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        set_run_font(r, "TH Sarabun New", 18, bold=True, color_rgb=RGBColor(30, 58, 138))
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        set_run_font(r, "TH Sarabun New", 15.5, bold=True, color_rgb=RGBColor(37, 99, 235))
        return p

    def add_body(text, bold_prefix="", indent=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if indent:
            p.paragraph_format.left_indent = Inches(0.25)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            set_run_font(r_pre, "TH Sarabun New", 14, bold=True, color_rgb=RGBColor(15, 23, 42))
        r_txt = p.add_run(text)
        set_run_font(r_txt, "TH Sarabun New", 14, color_rgb=RGBColor(51, 65, 85))
        return p

    def add_bullet(bold_label, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r_lbl = p.add_run(bold_label)
        set_run_font(r_lbl, "TH Sarabun New", 14, bold=True, color_rgb=RGBColor(30, 41, 59))
        r_txt = p.add_run(text)
        set_run_font(r_txt, "TH Sarabun New", 14, color_rgb=RGBColor(51, 65, 85))
        return p

    def add_callout_box(title, content_list, border_hex="2563EB", bg_hex="F8FAFC"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(6.5)
        cell = table.cell(0, 0)
        set_cell_background(cell, bg_hex)
        set_callout_border(cell, border_hex, "32")
        set_cell_margins(cell, top=130, bottom=130, left=180, right=180)

        p0 = cell.paragraphs[0]
        p0.paragraph_format.space_after = Pt(4)
        r0 = p0.add_run(title + "\n")
        set_run_font(r0, "TH Sarabun New", 14.5, bold=True, color_rgb=RGBColor(30, 58, 138))

        for idx, (label, val) in enumerate(content_list):
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if label:
                rl = p.add_run(label + ": ")
                set_run_font(rl, "TH Sarabun New", 13.5, bold=True, color_rgb=RGBColor(180, 83, 9) if "ภาษาคน" in label or "ภาพจำ" in label else RGBColor(13, 148, 136))
            rv = p.add_run(val)
            set_run_font(rv, "TH Sarabun New", 13.5, color_rgb=RGBColor(30, 41, 59))

        p_end = doc.add_paragraph()
        p_end.paragraph_format.space_before = Pt(2)
        p_end.paragraph_format.space_after = Pt(4)

    # ------------------ SECTION 1 ------------------
    add_h1("1. ภาพรวมและบริบทของผู้ก่อตั้ง (Overview & Historical Context)")
    add_body(" การให้คำปรึกษาแบบจิตวิเคราะห์ (Psychoanalytic Therapy) ถือเป็นต้นแบบ (Prototype) ของการบำบัดและการให้คำปรึกษาในปัจจุบัน จุดประกายและกระตุ้นเร้าให้เกิดการพัฒนาทฤษฎีอื่น ๆ ตามมาอย่างกว้างขวาง", indent=True)
    add_bullet("ผู้ก่อตั้ง: ", "ซิกมันด์ ฟรอยด์ (Sigmund Freud, 1856–1939) แพทย์และนักจิตวิทยาชาวออสเตรีย")
    add_bullet("อิทธิพลทางความคิด: ", "ฟรอยด์เติบโตในยุคที่วิทยาศาสตร์และฟิสิกส์กลศาสตร์กำลังเฟื่องฟู ทำให้เขามองระบบจิตใจของมนุษย์เหมือนระบบเครื่องจักรกลพลังงาน (Energy System) กล่าวคือ พลังงานทางจิตมีปริมาณคงที่ ไม่สูญหาย แต่สามารถถ่ายเท กักเก็บ หรือเปลี่ยนรูปได้")

    add_callout_box("💡 ถอดรหัสบริบทจิตวิเคราะห์ (เข้าใจง่ายสไตล์ภาษาคน)", [
        ("เปรียบเหมือน", "แว่นตานักโบราณคดี ที่เชื่อว่าพฤติกรรมแปลก ๆ ในปัจจุบันไม่ได้เกิดขึ้นลอย ๆ แต่มีรากเหง้ามาจาก 'ของเก่าที่ถูกฝังไว้ใต้ดิน' (จิตไร้สำนึกและบาดแผลวัยเด็ก)"),
        ("หัวใจสำคัญ", "การทำให้ 'จิตไร้สำนึก' (Unconscious) โผล่ขึ้นมาสู่ 'จิตสำนึก' (Conscious) เพื่อให้เจ้าตัวเข้าใจความจริงและหลุดพ้นจากอดีต")
    ], border_hex="3B82F6", bg_hex="EFF6FF")

    # ------------------ SECTION 2 ------------------
    add_h1("2. ทัศนะต่อธรรมชาติของมนุษย์และพลังทางจิต (Human Nature & Psychic Drives)")
    add_body("ฟรอยด์มองธรรมชาติของมนุษย์ในแง่ลบ (Pessimistic & Deterministic) โดยเชื่อว่ามนุษย์เดิมทีเป็นสิ่งมีชีวิตที่ไม่มีเหตุผล ไม่มีการขัดเกลาทางสังคม แต่ขับเคลื่อนด้วยแรงขับสัญชาตญาณดิบเพื่อตอบสนองความสุขพึงพอใจของตนเองเป็นหลัก (Pleasure Principle)", indent=True)
    
    add_h2("2.1 แรงขับทางสัญชาตญาณ 2 ประเภท (Basic Instincts)")
    add_bullet("1) สัญชาตญาณแห่งชีวิต (Eros / Life Instinct): ", "มุ่งเน้นการเอาชีวิตรอด การปกป้องตนเอง และการแสวงหาความพึงพอใจทางเพศ (ซึ่งฟรอยด์นิยามคำว่า 'เพศ' ไว้อย่างกว้างขวาง ครอบคลุมถึงความสุขทางร่างกายและความรัก)")
    add_bullet("2) สัญชาตญาณแห่งความตาย (Thanatos / Death Instinct): ", "แสดงออกในรูปของพลังทำลายล้าง ความก้าวร้าว ความรุนแรง และการทำร้ายตนเองหรือผู้อื่น")

    add_h2("2.2 พลังลิบิโด (Libido) และการรักษาสมดุลทางจิต")
    add_bullet("พลังลิบิโด (Libido): ", "คือพลังงานทางจิต (Psychic Energy) ที่ขับเคลื่อนสัญชาตญาณ มีปริมาณคงที่ ไม่สูญหายแต่เปลี่ยนสภาพได้")
    add_bullet("กลไกความตึงเครียด: ", "เมื่อสัญชาตญาณทางเพศหรือความก้าวร้าวถูกกระตุ้นเร้า จะเกิดความตึงเครียด (Tension) ในจิตใจ พลังลิบิโดจะผลักดันให้ร่างกายและจิตใจหาทางระบายพลังงานออกไปสู่เป้าหมายภายนอก เพื่อคลายความเครียดและคืนสู่สภาวะสมดุล")

    # ------------------ SECTION 3 ------------------
    add_h1("3. โครงสร้างบุคลิกภาพ (Personality Structure: Id, Ego, Superego)")
    add_body("ฟรอยด์อธิบายว่า บุคลิกภาพของมนุษย์ประกอบด้วยองค์ประกอบ 3 ส่วน ที่ต้องทำหน้าที่ประสานงานกันตลอดเวลา หากส่วนใดส่วนหนึ่งเสียสมดุล จะส่งผลให้เกิดปัญหาทางบุคลิกภาพและสุขภาพจิต", indent=True)

    # Table comparing Id, Ego, Superego
    t_struct = doc.add_table(rows=4, cols=4)
    t_struct.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_struct.autofit = False
    set_table_borders(t_struct)

    col_widths = [Inches(1.2), Inches(1.2), Inches(1.5), Inches(2.6)]
    for row in t_struct.rows:
        for idx, w in enumerate(col_widths):
            row.cells[idx].width = w

    headers_struct = ["โครงสร้าง", "ช่วงเวลาที่เกิด", "หลักการทำงาน", "หน้าที่และบทบาทสำคัญ"]
    for idx, text in enumerate(headers_struct):
        cell = t_struct.cell(0, idx)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, "TH Sarabun New", 13.5, bold=True, color_rgb=RGBColor(255, 255, 255))

    struct_data = [
        ("อิด (Id)", "มีมาแต่กำเนิด\n(ติดตัวมาตั้งแต่เกิด)", "Pleasure Principle\n(หลักความพึงพอใจ)", "แหล่งพลังงานจิตดิบ สัญชาตญาณเพศและความก้าวร้าว มุ่งระบายพลังทันทีผ่านการกระทำหรือจินตนาการเพื่อลดความตึงเครียด ไม่คำนึงถึงเหตุผลหรือกฎเกณฑ์ใด ๆ"),
        ("อีโก้ (Ego)", "พัฒนาขึ้นภายหลัง\n(เรียนรู้จากโลกจริง)", "Reality Principle\n(หลักความเป็นจริง)", "ทำหน้าที่เป็น 'คนกลาง/ผู้จัดการ' ไกล่เกลี่ยระหว่างความต้องการของ Id กับสภาพความเป็นจริงของสังคม ชะลอการระบายพลังงานของ Id ไว้จนกว่าจะพบเป้าหมายที่ปลอดภัย"),
        ("ซูเปอร์อีโก้\n(Superego)", "พัฒนาจากการอบรม\n(ซึมซับจากพ่อแม่/สังคม)", "Perfection Principle\n(หลักอุดมคติ/ศีลธรรม)", "ตัวแทนมโนธรรมสำนึกผิดชอบชั่วดี กฎหมาย และบรรทัดฐานสังคม คอยกดดันและยับยั้ง Id พร้อมกระตุ้นให้บุคคลทำสิ่งที่สมบูรณ์แบบตามอุดมคติมากกว่าความสุข")
    ]

    for row_idx, data in enumerate(struct_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = t_struct.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 90, 90, 110, 110)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            set_run_font(r, "TH Sarabun New", 13, bold=(col_idx==0), color_rgb=RGBColor(30, 41, 59))
            if col_idx in [0, 1, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(4)
    p_sp2.paragraph_format.space_after = Pt(2)

    add_callout_box("🎭 ภาพจำ 3 ฝ่ายในชีวิตประจำวัน", [
        ("Id (เด็กเอาแต่ใจ)", "'จะกินชานมไข่มุกเดี๋ยวนี้! จะโดดเรียนไปเล่นเกม!' (เรียกร้องตามสัญชาตญาณ)"),
        ("Superego (แม่ชีเจ้าระเบียบ)", "'ทำแบบนั้นชั่วร้าย ตกนรก เป็นบาป โดดเรียนคือคนไม่ดี!' (กดดันด้วยศีลธรรม)"),
        ("Ego (ผู้จัดการคนกลาง)", "'ตั้งใจเรียนให้เสร็จ 2 คาบนี้ก่อน แล้วตอนเย็นค่อยไปซื้อชานมกินเป็นรางวัล' (หาทางออกที่พอดีและทำได้จริง)"),
        ("จุดเน้นจิตวิทยา", "สุขภาพจิตที่ดีเกิดจาก Ego ที่แข็งแรง สามารถประนีประนอมความต้องการของ Id และ Superego ให้สอดคล้องกับโลกความจริงได้")
    ], border_hex="F59E0B", bg_hex="FFFBEB")

    # ------------------ SECTION 4 ------------------
    add_h1("4. กลไกการปกป้องตนเองของอีโก้ (Ego Defense Mechanisms)")
    add_body("เมื่ออีโก้ (Ego) ถูกคุกคามจากแรงผลักของ Id หรือการลงโทษของ Superego บุคคลจะเกิด 'ความวิตกกังวล' (Anxiety) อีโก้จึงต้องสร้างกลไกป้องกันตนเองขึ้นมาเพื่อลดความตึงเครียด โดยกลไกเหล่านี้ทำงานในระดับจิตไร้สำนึก (Unconscious) มีทั้งหมด 7 กลไกสำคัญ ดังนี้:", indent=True)

    def add_defense_item(num_title, quote_def, plain_exp, example_text):
        add_h2(num_title)
        add_body(quote_def, bold_prefix="• นิยามวิชาการ: ")
        add_body(plain_exp, bold_prefix="• ถอดรหัสภาษาคน: ")
        add_body(example_text, bold_prefix="• ตัวอย่างพฤติกรรม: ")

    add_defense_item(
        "4.1 การเก็บกด (Repression)",
        "การบีบ กัก บังคับ ความทรงจำ ความคิด และการรับรู้ที่คุกคามความรู้สึก ให้พ้นออกจากจิตรู้สำนึก ปรากฏในรูปการป้องกันมิให้มองเห็นความจริง หรือบิดเบือนความจริงเพื่อปกป้องอีโก้",
        "การกดเรื่องที่เจ็บปวด ปมชอกช้ำ หรือความรู้สึกผิด ลงไปซ่อนไว้ใต้พรมของจิตใจ ลืมไปเลยเหมือนไม่เคยเกิดขึ้น แต่มันยังไม่ตายและพร้อมระเบิดออกมาเป็นอาการทางจิต",
        "เด็กที่เคยถูกล่วงละเมิดในวัยเด็กจำเหตุการณ์นั้นไม่ได้เลย หรือการ 'ลืม' วันนัดตรวจหมอฟันที่ตนเองกลัวมาก"
    )

    add_defense_item(
        "4.2 การฉายภาพ / โยนบาป (Projection)",
        "การป้องกันตนเองเมื่อถูกคุกคามจากพลังของอิดและซูเปอร์อีโก้ โดยอีโก้หาทางออกด้วยการย้ายพลังความรู้สึกนั้นไปสู่บุคคลภายนอก",
        "ตัวเองมีความรู้สึกไม่ดีหรือข้อเสียบางอย่าง แต่รับตัวเองไม่ได้ จึงผลักความผิดหรือโยนความรู้สึกนั้นไปให้คนอื่นว่าเป็นคนคิด/ทำ",
        "นาย ก เกลียด นาย ข แต่ยอมรับไม่ได้ว่าตนเองมีใจอคติ จึงพูดว่า 'ข มันไม่ชอบฉัน มันคอยจ้องจะแกล้งฉัน' หรือคนขี้โกงที่มักระแวงว่าคนอื่นจ้องจะโกงตน"
    )

    add_defense_item(
        "4.3 การกระทำตรงกันข้าม (Reaction Formation)",
        "กลไกที่อีโก้ปรับเปลี่ยนสัญชาตญาณที่คุกคามให้แสดงออกเป็นพฤติกรรมในขั้วตรงข้ามอย่างมากเกินไปในรูปสลับด้าน เพื่อป้องกันการแสดงออกของความรู้สึกแท้จริง",
        "ข้างในรู้สึกแบบหนึ่ง (เช่น เกลียด กลัว อิจฉา) แต่แกล้งแสดงออกในขั้วตรงข้ามแบบ 'โอเวอร์แอ็กติ้ง' เพื่อกลบเกลื่อนความจริงไม่ให้ใครจับได้",
        "เกลียดเพื่อนร่วมงานคนนี้เข้าไส้ แต่พอเจอหน้ากลับวิ่งเข้าไปกอด พูดจาหวานเจี๊ยบ เอาอกเอาใจจนผิดปกติ หรือคนที่มีความต้องการทางเพศสูงมากแต่กลับทำตัวเป็นผู้นำต่อต้านสื่อลามกอย่างรุนแรง"
    )

    add_defense_item(
        "4.4 การติดตรึง (Fixation)",
        "สภาวะหยุดชะงัก ไม่สามารถก้าวไปสู่ขั้นของพัฒนาการที่สูงขึ้นได้ เพราะวิตกกังวลเมื่อต้องเผชิญกับสถานการณ์ใหม่ กลัวล้มเหลว กลัวไม่ได้รับการยอมรับ หรือกลัวการพลัดพรากจากสภาพเดิม",
        "ไม่ยอมโต ไม่ยอมก้าวไปข้างหน้า เพราะรู้สึกว่าที่เดิมปลอดภัยกว่า ยึดติดอยู่กับความสบายหรือความเคยชินเดิม ๆ",
        "วัยรุ่นที่ไม่ยอมออกจากบ้านไปเรียนมหาวิทยาลัย หรือผู้ใหญ่ที่ไม่กล้ารับตำแหน่งบริหารเพราะกลัวความรับผิดชอบ จึงขอทำงานตำแหน่งเดิมตลอดไป"
    )

    add_defense_item(
        "4.5 การถดถอย (Regression)",
        "การก้าวถอยหลังสู่พัฒนาการขั้นที่ต่ำกว่าเดิม โดยพฤติกรรมที่แสดงออกมักมีลักษณะเหมือนพฤติกรรมของเด็ก เพื่อเรียกคืนความรู้สึกมั่นคงปลอดภัย",
        "เวลาเจอปัญหาเครียด ๆ รับมือไม่ไหว เลยถอยหลังกลับไปทำตัวเป็นเด็ก เช่น ร้องไห้งอแง ปาข้าวของ หรือหนีกลับไปซุกอกพ่อแม่",
        "นักศึกษาปี 1 ปรับตัวเข้ากับมหาวิทยาลัยไม่ได้ ต้องนั่งรถกลับบ้านไปนอนกับแม่ทุกสุดสัปดาห์ หรือลาออกจากมหาวิทยาลัยแทนที่จะสู้ทน"
    )

    add_defense_item(
        "4.6 การเลียนแบบ (Identification)",
        "การรับเอาคุณลักษณะของบุคคลอื่นเข้าสู่บุคลิกภาพของตนเอง มีความสำคัญยิ่งต่อพัฒนาการบุคลิกภาพ แบ่งออกเป็น 4 ชนิดย่อย:",
        "การ 'ก๊อปปี้' นิสัย ค่านิยม หรือพฤติกรรมของคนอื่นมาใส่ตัวเอง เพื่อเสริมสร้างความมั่นคงหรือทดแทนสิ่งที่ตนขาดหายไป",
        "แบ่งเป็น 4 รูปแบบสำคัญที่ต้องจำ:"
    )

    # 4 sub-types of identification
    add_bullet("1) การเลียนแบบชนิดหลงรักตนเอง (Narcissistic Identification): ", "รับเอาคุณลักษณะของคนอื่นที่มีนิสัยหรือความสามารถตรงกันกับตนเข้ามา เช่น มีพรสวรรค์ด้านดนตรีหรือกีฬาเหมือนกัน จึงรู้สึกภูมิใจและเลียนแบบเขาเพราะเหมือนตนเอง")
    add_bullet("2) การเลียนแบบชนิดเป้าหมาย (Goal-oriented Identification): ", "เลียนแบบบุคคลที่มีคุณลักษณะที่ตนใฝ่ฝันอยากจะมี เช่น เด็กเล็กอยากแข็งแรงเหมือนพ่อ จึงเลียนแบบท่าทางและการออกกำลังกายของพ่อ")
    add_bullet("3) การเลียนแบบเชิงสูญเสีย (Object Loss Identification): ", "เลียนแบบเพื่อดึงสิ่งที่สูญเสียไปกลับคืนมา เช่น เด็กพยายามเอาชนะใจพ่อแม่ที่หย่าร้างหรือห่างเหิน โดยยอมทำตามกฎระเบียบและค่านิยมของพ่อแม่ทุกอย่างเพื่อแลกกับความรัก")
    add_bullet("4) การเลียนแบบเชิงอำนาจ (Identification with Authority / Aggressor): ", "รับเอาข้อห้าม คำสั่ง และกฎเกณฑ์ของพ่อแม่หรือผู้มีอำนาจเข้ามาไว้ในใจตนเอง **เป็นกลไกสำคัญที่สุดที่นำไปสู่การพัฒนามโนสำนึกผิดชอบชั่วดี (Superego)**")

    add_defense_item(
        "4.7 การย้ายที่ (Displacement)",
        "กระบวนการถ่ายเทพลังงานทางจิตที่เป็นสัญชาตญาณดั้งเดิมจากวัตถุเป้าหมายหนึ่งไปสู่สิ่งอื่นแทน โดยพลังงานและเป้าหมายคงเดิม แต่วัตถุเป้าหมายเปลี่ยนไป",
        "การเปลี่ยนที่ระบายอารมณ์/ความต้องการ จากสิ่งที่ทำไม่ได้หรืออันตราย ไปลงกับสิ่งของหรือคนที่ปลอดภัยกว่า (สำนวน 'เตะหมา')",
        "โกรธหัวหน้าแต่ด่าไม่ได้ เลยกลับบ้านไปเตะหมาหรือพาลใส่ลูก หรือในแง่พัฒนาการ: พลังดูดนมแม่ในวัยทารก ถูกย้ายที่เปลี่ยนเป็น การอมลูกอม สูบบุหรี่ กินจุกจิก การพูดเก่ง จนพัฒนาเป็น 'นักพูด'"
    )

    # Summary table for Defense Mechanisms
    add_h2("📊 ตารางสรุป 7 กลไกป้องกันตนเอง (Quick Review Table)")
    t_def = doc.add_table(rows=8, cols=3)
    t_def.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_def.autofit = False
    set_table_borders(t_def)

    widths_def = [Inches(1.8), Inches(2.2), Inches(2.5)]
    for row in t_def.rows:
        for idx, w in enumerate(widths_def):
            row.cells[idx].width = w

    headers_def = ["กลไกป้องกันตนเอง", "คีย์เวิร์ดจำง่าย (ภาษาคน)", "ตัวอย่างออกสอบบ่อย"]
    for idx, text in enumerate(headers_def):
        cell = t_def.cell(0, idx)
        set_cell_background(cell, "0F766E")
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, "TH Sarabun New", 13.5, bold=True, color_rgb=RGBColor(255, 255, 255))

    def_table_data = [
        ("1. การเก็บกด (Repression)", "กดลงใต้พรม, แกล้งลืมเรื่องเจ็บปวด", "จำเหตุการณ์ถูกทำร้ายในวัยเด็กไม่ได้"),
        ("2. การฉายภาพ (Projection)", "โยนบาป, ตัวเองเป็นแต่ด่าคนอื่น", "เกลียดเพื่อนแต่บอกว่าเพื่อนเกลียดตน"),
        ("3. การกระทำตรงกันข้าม (Reaction Formation)", "ข้างในเกลียด ข้างนอกหวานเจี๊ยบ", "เกลียดมากแต่เข้าไปกอดและยิ้มประจบ"),
        ("4. การติดตรึง (Fixation)", "หยุดชะงัก ไม่ยอมโต ไม่ก้าวต่อ", "กลัวล้มเหลวจึงไม่ยอมรับงานใหม่ที่ท้าทาย"),
        ("5. การถดถอย (Regression)", "ถอยหลังเป็นเด็ก, หนีกลับบ้าน", "เด็กโตกลับไปฉี่รดที่นอน/นศ.ลาออกหนีปัญหา"),
        ("6. การเลียนแบบ (Identification)", "ก๊อปปี้คนอื่น (4 แบบย่อย)", "เลียนแบบอำนาจพ่อแม่จนสร้าง Superego"),
        ("7. การย้ายที่ (Displacement)", "เตะหมา, เปลี่ยนที่ระบายพลังงาน", "โกรธครูแต่ไปพังโต๊ะ / ดูดนมแม่สู่นักพูด")
    ]

    for row_idx, data in enumerate(def_table_data, start=1):
        bg = "F0FDFA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = t_def.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            set_run_font(r, "TH Sarabun New", 13, bold=(col_idx==0), color_rgb=RGBColor(15, 23, 42))
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p_sp3 = doc.add_paragraph()
    p_sp3.paragraph_format.space_before = Pt(4)
    p_sp3.paragraph_format.space_after = Pt(2)

    # ------------------ SECTION 5 ------------------
    add_h1("5. เทคนิคและวิธีการให้คำปรึกษาทางจิตวิเคราะห์ (Counseling Techniques)")
    add_body("เป้าหมายสูงสุดของจิตวิเคราะห์ คือ การช่วยให้ผู้รับการปรึกษาเกิด 'ความเข้าใจตนเองอย่างลึกซึ้ง' (Insight) โดยการดึงสิ่งที่อยู่ในจิตไร้สำนึกให้ขึ้นมาสู่จิตสำนึก ผ่าน 6 เทคนิคหลัก:", indent=True)

    add_h2("5.1 การเชื่อมโยงอิสระ (Free Association)")
    add_bullet("หลักการ: ", "ขอให้ผู้รับการปรึกษาพูดหรือระบายทุกความคิด ความรู้สึก ความทรงจำที่ผุดขึ้นมาในจิตใจ โดยไม่ต้องเรียบเรียง ไม่ต้องมีเหตุผล และไม่ต้องกังวลว่าจะปะติดปะต่อกันหรือไม่")
    add_bullet("การจัดบรรยากาศ: ", "ผู้ให้การปรึกษาจะไม่นั่งประจันหน้า (มักให้นอนบนเตียง/โซฟา หรือนั่งเยื้องอยู่นอกสายตา) เพื่อลดความกดดันและความประหม่า พร้อมทั้งบันทึกข้อมูลอย่างละเอียดโดยไม่ขัดจังหวะ")

    add_h2("5.2 การระบายอารมณ์ (Catharsis / Abreaction)")
    add_bullet("หลักการ: ", "การช่วยให้ผู้รับการปรึกษารื้อฟื้นอารมณ์ความรู้สึกที่เคยเกิดขึ้นหรือถูกเก็บกดไว้ในอดีต โดยนำพาตนเองให้จินตนาการหรือย้อนกลับไปยังเหตุการณ์ต้นตอ เพื่อปลดปล่อยพลังงานทางจิตที่คั่งค้างออกมา")

    add_h2("5.3 การวิเคราะห์การต่อต้าน (Analysis of Resistance)")
    add_bullet("ธรรมชาติของการต่อต้าน: ", "เป็นสภาวะธรรมชาติที่เกิดขึ้นเสมอในสัมพันธภาพการปรึกษา เช่น การเงียบเป็นเวลานาน การเปลี่ยนเรื่องคุย การมาสาย หรือการไม่มาตามนัดหมาย")
    add_bullet("วิธีจัดการ: ", "ผู้ให้การปรึกษาต้องแสดงความเข้าใจและยอมรับในเบื้องต้นก่อน (ไม่ต่อว่า) จากนั้นจึงสนับสนุนให้ผู้รับการปรึกษาร่วมสำรวจว่าพฤติกรรมการต่อต้านนั้นมีความหมายหรือสะท้อนความเจ็บปวดใดในอดีต")

    add_h2("5.4 การวิเคราะห์การถ่ายโยงความรู้สึก (Analysis of Transference) ⭐ [สำคัญมาก]")
    add_bullet("นิยาม: ", "ภาวะที่ผู้รับการปรึกษานำเอาอารมณ์ ความรัก ความโกรธ หรือปมความขัดแย้งที่ยังค้างคากับบุคคลสำคัญในอดีต (เช่น พ่อ แม่) มา 'ถ่ายโยง' หรือสวมทับใส่ตัวผู้ให้การปรึกษาในปัจจุบัน")
    add_bullet("ขั้นตอนการจัดการของครู/ผู้ให้คำปรึกษา: ", "")
    add_body("1) สำรวจตนเองก่อนเสมอ: ตรวจสอบว่าตนเองได้พูดหรือกระทำสิ่งใดที่ไปกระตุ้นอารมณ์เด็กหรือไม่ หากไม่ใช่ แสดงว่าเป็นปฏิกิริยาถ่ายโยงของเด็กเอง", indent=True)
    add_body("2) ตอบสนองด้วยความเข้าใจ: ไม่โต้ตอบด้วยอารมณ์ ไม่โกรธตอบ แต่ยอมรับความรู้สึกนั้น", indent=True)
    add_body("3) ใช้คำถามเชื่อมโยงกลับสู่อดีต: ถามเชิญชวนให้พูดถึงความรู้สึกนั้น และสืบสาวกลับไปยังบุคคลในอดีตที่เด็กเคยมีความรู้สึกเช่นนี้ให้ไกลที่สุด", indent=True)

    add_h2("5.5 การวิเคราะห์ความฝัน (Dream Analysis)")
    add_bullet("หลักการ: ", "ยามหลับ กลไกการป้องกันตนจะผ่อนคลายลง ทำให้ความต้องการในจิตไร้สำนึกเล็ดลอดออกมาในรูปของสัญลักษณ์ความฝัน")
    add_bullet("เนื้อหา 2 ระดับในความฝัน: ", "")
    add_body("• เนื้อหาที่แสดงออก (Manifest Content): เรื่องราวและภาพความฝันที่บุคคลจำได้เมื่อตื่นนอน", indent=True)
    add_body("• เนื้อหาซ่อนเร้น (Latent Content): ความหมายที่แท้จริงและแรงขับทางสัญชาตญาณใต้จิตไร้สำนึกที่ต้องผ่านการตีความ (ผู้ใช้ต้องผ่านการฝึกอบรมเฉพาะทาง)", indent=True)

    add_h2("5.6 การแปลความ (Interpretation) และกฎเหล็ก 4 ข้อ")
    add_bullet("นิยาม: ", "การชี้แนะ อธิบายความหมายจากคำพูด ความฝัน การต่อต้าน และการถ่ายโยง เพื่อช่วยให้ผู้รับการปรึกษาตระหนักรู้และเข้าใจปัญหาของตนเอง (Insight)")
    add_bullet("⭐ ข้อควรคำนึง 4 ข้อในการแปลความ (ออกสอบบ่อย): ", "")
    add_body("1. สัมพันธภาพต้องพร้อม: ใช้เทคนิคนี้เมื่อมีความไว้วางใจและสัมพันธภาพแน่นแฟ้นแล้วเท่านั้น", indent=True)
    add_body("2. เริ่มจากเรื่องง่ายก่อน: ควรเริ่มแปลจากประเด็นที่ไม่ลึกซึ้งหรือง่ายต่อการเข้าใจก่อน อย่าเพิ่งแปลเรื่องลึกเกินไป", indent=True)
    add_body("3. ให้ผู้รับการปรึกษาแปลเองก่อน: ลองเปิดโอกาสให้ผู้รับการปรึกษาได้ฝึกแปลความหมายด้วยตนเอง", indent=True)
    add_body("4. ถ้าเขาปฏิเสธ ให้รั้งรอไว้: หากผู้รับการปรึกษาไม่ยอมรับการแปลความ ควรรั้งรอการแปลในประเด็นอื่นไว้ก่อน อย่าเพิ่งฝืนยัดเยียด", indent=True)

    # ------------------ SECTION 6 ------------------
    add_h1("6. บทบาทครูกับการประยุกต์ใช้ทฤษฎีจิตวิเคราะห์ในโรงเรียน")
    add_body("แม้ครูจะไม่ใช่นักจิตวิเคราะห์ที่ต้องนำเด็กมานอนบนโซฟา แต่แนวคิดของฟรอยด์มีประโยชน์อย่างยิ่งต่อการเข้าใจพฤติกรรมของนักเรียน:", indent=True)
    add_bullet("1. มองทะลุพฤติกรรมก้าวร้าว: ", "เข้าใจว่าเด็กที่ชอบเกเร แกล้งเพื่อน หรือดื้อต่อต้าน อาจกำลังใช้กลไก Displacement (ย้ายที่ระบายอารมณ์จากการถูกพ่อแม่ทำร้าย) หรือ Projection (ฉายภาพความกลัว)")
    add_bullet("2. รับมือกับ Transference: ", "เมื่อนักเรียนแสดงความเกลียดชังหรือไม่ชอบครูโดยไม่มีเหตุผล ครูต้องตั้งสติ ไม่โกรธตอบ เพราะรู้ว่าตนเองกำลังเป็น 'ตัวแทนของพ่อแม่' ที่เด็กกำลังถ่ายโยงอารมณ์มาใส่")
    add_bullet("3. เสริมความเข้มแข็งให้ Ego: ", "ช่วยเหลือนักเรียนให้เผชิญหน้ากับความจริง ฝึกทักษะการตัดสินใจและการแก้ปัญหา แทนที่จะปล่อยให้เด็กใช้กลไกถดถอย (Regression) หรือหนีปัญหา")

    # ------------------ SECTION 7 ------------------
    add_h1("7. สรุปประเด็นเน้นจำสำหรับเตรียมสอบ (Exam High-Yield Summary)")

    add_callout_box("🎯 Checklists 10 ประเด็นที่มักออกข้อสอบประจำบทที่ 8", [
        ("1. ผู้ก่อตั้งและรากคิด", "Sigmund Freud ชาวออสเตรีย มองพลังจิตคล้ายพลังงานกลศาสตร์ (ไม่สูญหายแต่เปลี่ยนรูปได้)"),
        ("2. มุมมองธรรมชาติมนุษย์", "มองในแง่ลบ (Pessimistic) เดิมไม่มีเหตุผล มุ่งตอบสนองสัญชาตญาณตนเอง (Pleasure Principle)"),
        ("3. สัญชาตญาณ 2 ขั้ว", "Eros (ชีวิต/เพศ/รอด) vs Thanatos (ความตาย/ก้าวร้าว/ทำลาย) ขับเคลื่อนด้วยพลัง Libido"),
        ("4. โครงสร้างจิต 3 ตัว", "Id = เกิดพร้อมเรา/สุขสำราญ, Ego = เกิดทีหลัง/ความจริง/คนกลาง, Superego = มโนธรรม/อุดมคติ"),
        ("5. อาการทางจิตเกิดจาก", "Ego อ่อนแอ ไม่สามารถไกล่เกลี่ยความต้องการระหว่าง Id กับ Superego ได้"),
        ("6. การสร้างมโนสำนึก", "เกิดจาก 'การเลียนแบบเชิงอำนาจ (Identification with Authority)' จากพ่อแม่"),
        ("7. Reaction Formation", "ข้างในคิดลบ แต่แสดงพฤติกรรมขั้วตรงข้ามอย่างโอเวอร์ (เกลียดแต่เข้าไปกอดประจบ)"),
        ("8. Transference", "ถ่ายโยงปมเก่าในอดีตมาใส่ผู้ให้คำปรึกษา -> ผู้ให้คำปรึกษาต้องสำรวจตนเองก่อนเป็นอันดับแรก"),
        ("9. ความฝัน 2 ระดับ", "Manifest Content (จำได้ตอนตื่น) vs Latent Content (ความหมายซ่อนเร้นแท้จริงใต้จิตไร้สำนึก)"),
        ("10. กฎการแปลความ", "สัมพันธภาพต้องดี เริ่มจากง่ายไปยาก ให้เด็กแปลเองก่อน ถ้าเด็กปฏิเสธให้หยุดรั้งรอไว้")
    ], border_hex="DC2626", bg_hex="FEF2F2")

    # Save document
    out_dir = r"C:\Users\chinj\OneDrive\ปี3\EDP3502-ช่วยเหลือ"
    out_filename = "สรุปเนื้อหา_EDP3502_บทที่8_การให้คำปรึกษาทฤษฎีแบบจิตวิเคราะห์.docx"
    out_path = os.path.join(out_dir, out_filename)
    doc.save(out_path)
    print(f"Document saved successfully at: {out_path}")

    # Also save a copy in 'งาน' folder for backup
    work_dir = os.path.join(out_dir, "งาน")
    if os.path.exists(work_dir):
        work_path = os.path.join(work_dir, out_filename)
        doc.save(work_path)
        print(f"Backup copy saved at: {work_path}")

if __name__ == "__main__":
    build_chapter8_doc()
