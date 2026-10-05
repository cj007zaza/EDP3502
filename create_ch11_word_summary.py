import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_callout_border(cell, border_hex="2563EB", border_size="32"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def create_summary_doc():
    doc = docx.Document()

    # Margins: 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base style
    normal = doc.styles['Normal']
    normal.font.name = 'TH Sarabun New'
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)

    def add_title_block():
        # Title table banner
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.5)
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "1E3A8A") # Navy Blue
        set_cell_margins(cell, top=260, bottom=260, left=260, right=260)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r0 = p.add_run("เอกสารสรุปเนื้อหารายวิชา EDP3502 การให้คำปรึกษาและช่วยเหลือผู้เรียน\n")
        r0.font.name = "TH Sarabun New"
        r0.font.size = Pt(14)
        r0.font.color.rgb = RGBColor(191, 219, 254) # Blue 200

        r1 = p.add_run("บทที่ 11: ทฤษฎีการให้คำปรึกษาแบบเกสตอลต์\n(Gestalt Counseling & Therapy)\n")
        r1.bold = True
        r1.font.name = "TH Sarabun New"
        r1.font.size = Pt(22)
        r1.font.color.rgb = RGBColor(255, 255, 255)

        r2 = p.add_run("ผู้ริเริ่มและพัฒนาทฤษฎี: ฟริตซ์ เพิร์ลส์ (Fritz Perls) | เอกสารสรุปฉบับสมบูรณ์")
        r2.italic = True
        r2.font.name = "TH Sarabun New"
        r2.font.size = Pt(13)
        r2.font.color.rgb = RGBColor(224, 231, 255)

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(8)
        p_spacer.paragraph_format.space_after = Pt(6)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'TH Sarabun New'
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(30, 58, 138) # Navy
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'TH Sarabun New'
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(37, 99, 235) # Royal Blue
        return p

    def add_body(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = 'TH Sarabun New'
            r_bold.font.size = Pt(14)
            r_bold.font.color.rgb = RGBColor(15, 23, 42)
        run = p.add_run(text)
        run.font.name = 'TH Sarabun New'
        run.font.size = Pt(14)
        run.italic = italic
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = 'TH Sarabun New'
            r_bold.font.size = Pt(14)
            r_bold.font.color.rgb = RGBColor(15, 23, 42)
        run = p.add_run(text)
        run.font.name = 'TH Sarabun New'
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_box(title, text, bg="F8FAFC", border="2563EB", icon="💡"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.5)
        cell = tbl.cell(0, 0)
        set_cell_background(cell, bg)
        set_callout_border(cell, border, "32")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(f"{icon} {title}\n")
        r_title.bold = True
        r_title.font.name = "TH Sarabun New"
        r_title.font.size = Pt(14.5)
        r_title.font.color.rgb = RGBColor(30, 58, 138)

        r_text = p.add_run(text)
        r_text.font.name = "TH Sarabun New"
        r_text.font.size = Pt(13.5)
        r_text.font.color.rgb = RGBColor(30, 41, 59)

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(2)
        p_spacer.paragraph_format.space_after = Pt(4)

    # ---------------- BUILD CONTENT ----------------
    add_title_block()

    # Section 1
    add_h1("1. บทนำและหัวใจสำคัญของทฤษฎีเกสตอลต์ (Gestalt Counseling Overview)")
    add_body(
        "คำว่า \"เกสตอลต์\" (Gestalt) เป็นคำในภาษาเยอรมัน มีความหมายตรงกับภาษาอังกฤษว่า Form, Pattern หรือ Configuration ซึ่งในทางจิตวิทยาหมายถึง "
        "\"ส่วนรวมทั้งหมดที่มีความหมาย\" (A meaningful whole) โดยมองว่าผลรวมของส่วนทั้งหมดมีความหมายมากกว่าการนำส่วนย่อยๆ แต่ละส่วนมารวมกัน "
        "ทฤษฎีนี้ริเริ่มและพัฒนาโดย ฟริตซ์ เพิร์ลส์ (Fritz Perls) โดยผสมผสานรากฐานทางปรัชญาสำคัญ 2 สำนัก ได้แก่:",
        bold_prefix="นิยามและความหมาย: "
    )
    add_bullet("ยึดการรับรู้และความเป็นจริงตามประสบการณ์ตรงเฉพาะบุคคล (Subjective / Phenomenological Reality) เกสตอลต์จะไม่ตัดสินว่าสิ่งใดคือความจริงสากลที่แท้จริง แต่ให้ความสำคัญกับความจริงตามที่บุคคลรับรู้และรู้สึกจากประสบการณ์ภายใน", bold_prefix="1) ปรากฏการณ์นิยม (Phenomenology): ")
    add_bullet("เน้นการมีชีวิตอยู่และตระหนักรู้กับ \"ที่นี่และเดี๋ยวนี้\" (Here and Now) รวมถึงส่งเสริมความรับผิดชอบ (Responsibility) เสรีภาพในการเลือก และการพึ่งพาตนเอง (Self-reliance)", bold_prefix="2) อัตถิภาวนิยม (Existentialism): ")

    add_box(
        "จุดเด่นที่แตกต่างจากทฤษฎีจิตวิเคราะห์เดิม (Gestalt vs Psychoanalysis)",
        "เกสตอลต์ปฏิเสธการสืบค้นหาสาเหตุในอดีต (ไม่ถามว่า \"ทำไม\" หรือ \"Why\") และไม่เน้นการวิเคราะห์ตีความโดยผู้ให้คำปรึกษา "
        "แต่หันมาให้ความสำคัญกับ \"อะไร\" และ \"อย่างไร\" (What & How) ที่กำลังเกิดขึ้นในสภาวะปัจจุบัน โดยเชื่อว่า "
        "\"การมีสติตระหนักรู้ (Awareness) มีธรรมชาติของการบำบัดรักษาในตัวเอง\" เมื่อบุคคลรู้ตัวอย่างแท้จริง เขาจะสามารถบูรณาการและเยียวยาตนเองได้",
        bg="EFF6FF", border="3B82F6", icon="🌟"
    )

    # Section 2
    add_h1("2. ธรรมชาติมนุษย์และแนวคิดพื้นฐาน (Key Concepts)")

    add_h2("2.1 ความสมดุลของอินทรีย์ (Organismic Self-Regulation & Homeostasis)")
    add_body(
        "พฤติกรรมทั้งมวลของมนุษย์ถูกผลักดันด้วยสัญชาตญาณในการพยายามรักษาสภาวะสมดุล เมื่อใดก็ตามที่อินทรีย์ (ทั้งทางร่างกายและจิตใจ) เกิดความต้องการ "
        "สภาวะสมดุลจะถูกรบกวน อินทรีย์จะมีแรงขับเคลื่อนเพื่อหาวิถีทางตอบสนองความต้องการนั้น เมื่อบรรลุเป้าหมายแล้ว อินทรีย์จะกลับคืนสู่สภาวะสมดุล "
        "ความต้องการเดิมจะสงบลง และเปิดโอกาสให้ความต้องการใหม่ๆ ทยอยก้าวเข้ามาหมุนเวียนสลับกัน กระบวนการรักษาสมดุลอย่างต่อเนื่องนี้เป็นหัวใจจำเป็นของการเติบโตงอกงามของมนุษย์"
    )

    add_h2("2.2 ความเป็นองค์รวม (Holism)")
    add_body(
        "มนุษย์คือองค์รวมที่ไม่สามารถแยกส่วนได้ การทำงานของระบบสรีระ (กาย) และระบบจิตใจ (จิต) ทำงานบูรณาการสัมพันธ์ซึ่งกันและกันอย่างแนบแน่น "
        "นอกจากนี้ มนุษย์ยังไม่สามารถแยกขาดออกจากสิ่งแวดล้อมได้ เพราะการดำรงชีวิตต้องพึ่งพาอากาศ น้ำ อาหาร และการมีปฏิสัมพันธ์กับผู้อื่นในสังคมตลอดเวลา"
    )

    add_h2("2.3 ความก้าวร้าวเชิงบวกในพัฒนาการ (Healthy Aggression)")
    add_body(
        "การรักษาสมดุลและการพัฒนาตนให้งอกงามจำเป็นต้องอาศัยความก้าวร้าวในระดับที่เหมาะสม ซึ่งเป็นพัฒนาการตามธรรมชาติตั้งแต่วัยทารก "
        "(เช่น ทารกต้องใช้ฟันขบเคี้ยวอาหารเพื่อย่อย หรือคายอาหารที่ทำให้ไม่สบายกายทิ้งออกมา) ความก้าวร้าวช่วยให้มนุษย์กล้าที่จะติดต่อสัมพันธ์ "
        "กล้าปฏิเสธสิ่งที่ไม่ดี และกล้าเลือกสิ่งที่มีประโยชน์เพื่อตอบสนองความต้องการตนเอง"
    )

    add_h2("2.4 มโนทัศน์เรื่อง \"ภาพ\" (Figure) และ \"พื้น\" (Ground)")
    add_body(
        "ในกระแสชีวิตที่เปลี่ยนแปลงตลอดเวลา บุคคลจะรับรู้สิ่งต่างๆ ผ่านกระบวนการจัดภาพและพื้น:",
        bold_prefix="หลักการทำงาน: "
    )
    add_bullet("คือสิ่งหรือความต้องการเด่นชัดที่บุคคลกำลังให้ความสนใจและมีสมาธิอยู่ในขณะนั้น มีขอบเขตชัดเจน โดดเด่นขึ้นมา", bold_prefix="ภาพ (Figure): ")
    add_bullet("คือบริบทแวดล้อม ข้อมูล หรือความต้องการอื่นๆ ที่บุคคลยังไม่ได้ให้ความสนใจในขณะนั้น ขอบเขตเลือนราง ไม่แน่นอน", bold_prefix="พื้น (Ground): ")
    add_body(
        "ตัวอย่างคลาสสิก: ขณะที่เรากำลังนั่งฟังเพลงโปรดแล้วมีเสียงฝนตกอยู่นอกหน้าต่าง หากเราตั้งใจฟังดนตรี เสียงเพลงจะเป็น \"ภาพ\" ส่วนเสียงฝนจะเป็น \"พื้น\" "
        "แต่เมื่อใดที่เราสะดุ้งหันไปฟังว่าฝนตกหนักหรือไม่ เสียงฝนจะพุ่งขึ้นมากลายเป็น \"ภาพ\" และเสียงเพลงจะถอยลงไปเป็น \"พื้น\" สลับสับเปลี่ยนกันตามความสนใจ",
        italic=True
    )

    # Section 3
    add_h1("3. นิยามของ \"สุขภาพจิต\" ตามแนวคิดเกสตอลต์ (Gestalt Mental Health)")
    add_body(
        "สุขภาพจิตตามมุมมองเกสตอลต์ หมายถึง การที่บุคคลมีส่วนรวมที่มีความหมายแบบรวมขั้ว มีภาพและพื้นที่ชัดเจน และมีการปรับตัวอย่างสร้างสรรค์ (Creative Adjustment) โดยจำแนกสภาวะได้ดังนี้:"
    )

    # Table of Mental Health
    t_mh = doc.add_table(rows=4, cols=2)
    t_mh.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_mh.autofit = False
    set_table_borders(t_mh, "CBD5E1", "4")
    col_w = [Inches(2.2), Inches(4.3)]
    for row in t_mh.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = col_w[idx]

    # Headers
    h_titles = ["สภาวะสุขภาพจิต", "ลักษณะของภาพ (Figure) และพื้น (Ground)"]
    for i, title in enumerate(h_titles):
        cell = t_mh.cell(0, i)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.name = "TH Sarabun New"
        r.font.size = Pt(14.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    mh_data = [
        ("ผู้มีสุขภาพจิตดี\n(Healthy State)", "ภาพมีความเด่นชัด แตกต่างจากพื้นอย่างมีขอบเขต และมีการสลับสับเปลี่ยนระหว่างความต้องการต่างๆ อย่างมีจังหวะเหมาะสม เมื่อได้รับการตอบสนอง ภาพจะกลับไปเป็นพื้น และเปิดโอกาสให้ความต้องการใหม่ก้าวขึ้นมาเป็นภาพอย่างยืดหยุ่น"),
        ("สภาวะบุคลิกภาพแปรปรวน / ฮิสทีเรีย (Histrionic)", "ความต้องการที่เข้ามาเป็นภาพสับเปลี่ยนรวดเร็วพรั่งพรูจนเกินไป เข้าและออกอย่างไม่เป็นระเบียบ ทำให้จิตใจว้าวุ่น ขาดความมั่นคงและจังหวะที่เหมาะสม"),
        ("สภาวะย้ำคิดย้ำทำ (Obsessive-Compulsive)", "ความต้องการที่เข้ามาเป็นภาพติดแน่น ไม่ยอมเคลื่อนที่กลับไปเป็นพื้น จนความต้องการอื่นแทบไม่มีช่องว่างเข้ามาแทนที่ เปรียบเหมือนการมองเห็นภาพที่ปราศจากพื้น"),
    ]

    for row_idx, (st, desc) in enumerate(mh_data, start=1):
        c0 = t_mh.cell(row_idx, 0)
        c1 = t_mh.cell(row_idx, 1)
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_margins(c0, top=100, bottom=100, left=140, right=140)
        set_cell_margins(c1, top=100, bottom=100, left=140, right=140)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(st)
        r0.bold = True
        r0.font.name = "TH Sarabun New"
        r0.font.size = Pt(13.5)
        r0.font.color.rgb = RGBColor(30, 58, 138)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(desc)
        r1.font.name = "TH Sarabun New"
        r1.font.size = Pt(13.5)
        r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_body(
        "นอกจากนี้ คนที่ไม่ทราบว่าตนเองต้องการอะไร สะท้อนถึงสภาวะที่ไม่ปรากฏภาพใดๆ ชัดเจนเลย ความต้องการทั้งหลายจมแฝงอยู่ในพื้น ทำให้เกสตอลต์ดูพร่ามัว "
        "และในแง่การปรับตัวอย่างสร้างสรรค์ บุคคลที่มีสุขภาพจิตดีต้องรักษาสมดุลระหว่างตนเองกับสิ่งแวดล้อม หากหมกมุ่นรักตนเองมากเกินไป (Narcissism) "
        "จนปฏิเสธโลกภายนอก ก็จัดเป็นสภาวะที่ขาดสมดุลเช่นกัน"
    )

    # Section 4
    add_h1("4. มโนทัศน์ด้านโครงสร้างและขอบเขตบุคลิกภาพ (Structure of Personality)")

    add_h2("4.1 ขั้วตรงกันข้ามในบุคลิกภาพ (Polarities)")
    add_body(
        "แม้ธรรมชาติมนุษย์จะเป็นองค์รวม แต่ในชีวิตจริงคนจำนวนมากมีบุคลิกภาพที่แตกแยกเป็นเสี่ยงๆ เนื่องจากบุคคลมักตระหนักรู้และยอมรับเฉพาะบุคลิกภาพส่วนที่ตนพึงใจ "
        "และผลักไส ปฏิเสธ หรือตัดขาดบุคลิกภาพส่วนที่อยู่ขั้วตรงข้าม เช่น คนที่พยายามทำตัวเข้มแข็งเป็นอิสระตลอดเวลา มักปฏิเสธความต้องการกำลังใจและความอ่อนแอในส่วนลึก "
        "การปฏิเสธขั้วตรงข้ามนี้เกิดขึ้นเพราะบุคคลกลัวความไม่สบายใจที่จะยอมรับจุดอ่อนตนเอง จึงฝืนสร้างสมดุลจอมปลอมขึ้นมา เกสตอลต์จึงมุ่งหมายช่วยให้บุคคล "
        "\"ยอมรับและบูรณาการขั้วตรงข้ามทั้งสองเข้าด้วยกัน\""
    )

    add_h2("4.2 ขอบเขตของอีโก้ (Ego Boundary)")
    add_body(
        "ขอบเขตของอีโก้คือเส้นแบ่งอาณาเขตที่กำหนดว่าสิ่งใดคือ \"ตนเอง\" (Self) และสิ่งใดคือ \"สิ่งแวดล้อมภายนอก\" (Not-Self) โดยทำงานผ่าน 2 กลไกสำคัญ:",
        bold_prefix="การทำงานของขอบเขต: "
    )
    add_bullet("การรับเอาค่านิยม ความเชื่อ อารมณ์ และความรู้สึกของผู้อื่น (เช่น พ่อแม่ คนรัก สังคม) เข้ามาเป็นส่วนหนึ่งในขอบเขตของตน", bold_prefix="การเลียนแบบ (Identification): ")
    add_bullet("การพิจารณาว่าสิ่งนั้น ความเชื่อนั้น หรืออารมณ์นั้น มิใช่ของตน และผลักดันออกไปอยู่นอกขอบเขตของอีโก้", bold_prefix="การแปลกแยก (Alienation): ")
    add_bullet("ขอบเขตอีโก้ที่อ่อนแอ/หละหลวม: จะรับเอาความคิดและความรู้สึกของคนอื่นเข้ามามากเกินไป จนแยกไม่ออกว่าตนเองต้องการอะไร ทึกทักว่าคนอื่นรู้สึกเหมือนตน", bold_prefix="ปัญหาที่พบ: ")
    add_bullet("ขอบเขตอีโก้ที่เข้มงวดเกินไป: จะปิดกั้นประสบการณ์ใหม่ๆ แข็งกระด้าง ไม่ยืดหยุ่น และรู้สึกแปลกแยกโดดเดี่ยวจากสังคม", bold_prefix="ปัญหาที่พบ: ")

    # Section 5
    add_h1("5. ลำดับขั้นของอาการประสาท / การติดต่อสัมพันธ์ 5 ระดับ (5 Layers of Neurosis)")
    add_body(
        "ในการเจริญเติบโตและการให้คำปรึกษา บุคคลจะต้องก้าวข้ามชั้นของการป้องกันตนเอง (Layers of Contact/Neurosis) ทั้ง 5 ระดับ เพื่อเข้าถึงตัวตนที่แท้จริงและเป็นองค์รวม ดังนี้:"
    )

    # 5 Layers Table
    t_layers = doc.add_table(rows=6, cols=3)
    t_layers.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_layers.autofit = False
    set_table_borders(t_layers, "CBD5E1", "4")
    l_w = [Inches(1.8), Inches(2.2), Inches(2.5)]
    for row in t_layers.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = l_w[idx]

    l_heads = ["ระดับ (Layer)", "สภาพจิตใจและพฤติกรรม", "ความหมายเชิงการบำบัด"]
    for i, title in enumerate(l_heads):
        cell = t_layers.cell(0, i)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.name = "TH Sarabun New"
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(255, 255, 255)

    layers_data = [
        ("1. ระดับเสแสร้ง\n(Phony Layer)", "ปฏิสัมพันธ์ตามธรรมเนียม มารยาท สวมบทบาทที่สังคมมอบหมาย ไม่เป็นตัวของตัวเองอย่างแท้จริง", "บุคคลแสดงไปตามบทบาทเพื่อปกปิดตัวตนจริงและเอาตัวรอดในสังคม"),
        ("2. ระดับกลัว\n(Phobic Layer)", "หวาดกลัวที่จะเผชิญหน้ากับความบกพร่องหรือส่วนที่ตนไม่ต้องการ กลัวการถูกปฏิเสธจากผู้อื่น", "เริ่มตระหนักถึงความเปราะบาง แต่พยายามปกปิดและหลีกเลี่ยงการเปิดเผยความรู้สึกแท้จริง"),
        ("3. ระดับทางตัน\n(Impasse Layer)", "รู้สึกหมดหนทาง ไร้พลังความสามารถภายใน รู้สึกว่างเปล่า ดำเนินชีวิตด้วยตนเองไม่ได้", "จุดวิกฤตที่บุคคลหยุดพึ่งตนเอง และเรียกร้องให้คนอื่นมาคิด ตัดสินใจ หรืออุ้มชูตนแทน"),
        ("4. ระดับระเบิดทางตัน\n(Implosive Layer)", "เริ่มมีสติตระหนักรู้ เผชิญกับความจริง ปลดล็อกเรื่องราวคั่งค้างในอดีต (Unfinished Business)", "ยอมรับความรู้สึกที่เคยปฏิเสธ ลดการป้องกันตัว เริ่มสัมผัสความเป็นไปได้ที่จะพึ่งพาตนเอง"),
        ("5. ระดับระเบิดสู่ภายนอก\n(Explosive Layer)", "ปลดปล่อยพลังงานที่ถูกกักขังไว้ แสดงอารมณ์ความรู้สึกอย่างแท้จริง กล้าแสดงออก มีชีวิตชีวา", "บุคคลเข้าถึงประสบการณ์การมีตนที่เป็นองค์รวม (Wholeness) ติดต่อสัมพันธ์กับโลกได้อย่างแท้จริง"),
    ]

    for row_idx, (col0, col1, col2) in enumerate(layers_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for c_idx, text_val in enumerate([col0, col1, col2]):
            c = t_layers.cell(row_idx, c_idx)
            set_cell_background(c, bg)
            set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            p = c.paragraphs[0]
            r = p.add_run(text_val)
            r.font.name = "TH Sarabun New"
            r.font.size = Pt(13)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(30, 58, 138)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Section 6
    add_h1("6. กลไกขัดขวางการติดต่อสัมพันธ์ / พฤติกรรมที่ไม่เหมาะสม (4 Boundary Disturbances)")
    add_body(
        "\"อาการประสาท\" ในมุมมองเกสตอลต์เกิดจากการที่บุคคลไม่อยู่กับปัจจุบัน ขาดการมีสติตระหนักรู้ และไม่สามารถบอกความต้องการของตนได้ "
        "ทำให้เรื่องราวในอดีตคั่งค้างไม่ได้รับการสะสาง (Unfinished Business) และใช้กลไกการแปลกแยกกับตนเอง 4 ประการเพื่อหลีกเลี่ยงการเผชิญความจริง:"
    )

    # 4 Disturbances Table
    t_dist = doc.add_table(rows=5, cols=4)
    t_dist.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_dist.autofit = False
    set_table_borders(t_dist, "CBD5E1", "4")
    d_w = [Inches(1.5), Inches(1.8), Inches(1.6), Inches(1.6)]
    for row in t_dist.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = d_w[idx]

    d_heads = ["กลไกทางจิต", "ลักษณะนิยาม", "พฤติกรรม / ตัวอย่าง", "แนวทางแก้ไขในห้องปรึกษา"]
    for i, title in enumerate(d_heads):
        cell = t_dist.cell(0, i)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=140, bottom=140, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.name = "TH Sarabun New"
        r.font.size = Pt(13.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    dist_data = [
        ("1. การกลืนค่านิยม\n(Introjection)", 
         "รับเอาความคิด ค่านิยม ความเชื่อเข้ามาทั้งหมดโดยไม่คิดไตร่ตรอง เปรียบเหมือนกลืนอาหารลงกระเพาะโดยไม่เคี้ยว", 
         "เป็นคนว่าง่าย รับทุกอย่างที่ผู้ให้คำปรึกษาบอกโดยไม่ขัด แต่ไม่รู้ว่าตัวเองเชื่ออะไรแท้จริง รู้สึกเหมือนเสแสร้ง", 
         "ระวังอย่าด่วนสรุปว่าการยอมทำตามคือการเปลี่ยนจริง ต้องกระตุ้นให้ผู้รับการปรึกษาฝึกคิดวิเคราะห์และเลือกสิ่งที่เหมาะกับตน"),
        ("2. การฉายภาพ\n(Projection)", 
         "ปฏิเสธคุณลักษณะหรือความรู้สึกที่แท้จริงในตนเอง แล้วโยนความรู้สึกผิด/ความรับผิดชอบนั้นไปให้ผู้อื่น", 
         "ก ต่อต้านและประณามว่านางสาว ข แต่งตัวยั่วยวน ทั้งที่ความจริงตนเองไม่ยอมรับความต้องการทางเพศของตน", 
         "ช่วยเหลือให้เกิดประสบการณ์การเป็นเจ้าของ (Owning) อารมณ์ ความรู้สึก และทัศนะที่แท้จริงของตนเอง"),
        ("3. การย้อนกลับสู่ตน\n(Retroflection)", 
         "กระทำต่อตนเองในสิ่งที่ความรู้สึกแท้จริงอยากกระทำต่อผู้อื่น ใช้ตนเองเป็นเป้าทดแทนสิ่งแวดล้อม", 
         "โกรธคนอื่นแต่หันมาทำร้ายตัวเอง กัดเล็บ กัดปาก กล่าวโทษตนเอง เกิดความเครียดและโรคทางกาย", 
         "ปลดปล่อยพลังงานที่ติดขัด ช่วยให้ผู้รับการปรึกษากล้าแสดงความรู้สึกและจัดการกับเป้าหมายภายนอกอย่างเหมาะสม"),
        ("4. ความหละหลวม\n(Confluence)", 
         "ไม่มีเส้นแบ่งอาณาเขตระหว่างตนกับสิ่งแวดล้อม รู้สึกเป็นเนื้อเดียวกัน ยอมรับความแตกต่างของคนอื่นไม่ได้", 
         "พ่อแม่คิดว่าลูกเป็นส่วนหนึ่งของตน หรือเพื่อนสนิทที่บังคับให้ทุกคนต้องคิดเหมือนกัน หากใครแยกตัวจะโกรธจัด", 
         "สอนให้รู้จักแยกตนเองออกจากผู้อื่น ตระหนักรู้ในความเป็นปัจเจกบุคคลที่มีเอกลักษณ์โดยไม่จำเป็นต้องรวมเป็นเนื้อเดียว"),
    ]

    for row_idx, row_vals in enumerate(dist_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            c = t_dist.cell(row_idx, c_idx)
            set_cell_background(c, bg)
            set_cell_margins(c, top=100, bottom=100, left=110, right=110)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "TH Sarabun New"
            r.font.size = Pt(12.5)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(30, 58, 138)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Section 7
    add_h1("7. เป้าหมายและบทบาทของผู้ให้คำปรึกษา (Goals & Counselor's Role)")

    add_h2("7.1 เป้าหมายของการปรึกษาแบบเกสตอลต์")
    add_bullet("มุ่งให้บุคคลมีสติรู้เท่าทันการสัมผัสผัสสะทั้ง 6 (รูป รส กลิ่น เสียง สัมผัส อารมณ์) และการกระทำของตนในขณะนั้น", bold_prefix="1) เพิ่มการมีสติตระหนักรู้ (Awareness): ")
    add_bullet("ช่วยให้บุคคลหยุดการคร่ำครวญถึงอดีต และหยุดความวิตกกังวลต่ออนาคต (ความกังวลคือช่องว่างระหว่างปัจจุบันกับอนาคต)", bold_prefix="2) ดำเนินชีวิตอยู่กับปัจจุบัน เดี๋ยวนี้ และที่นี่ (Here & Now): ")
    add_bullet("สะสางความรู้สึกโกรธ ความคับแค้นใจ หรือความเศร้าที่คั่งค้างให้สำเร็จบริบูรณ์ เพื่อคืนพลังงานชีวิตกลับมา", bold_prefix="3) จัดการเรื่องราวในอดีตที่คั่งค้าง (Unfinished Business): ")
    add_bullet("เปลี่ยนผ่านจากการร้องขอการประคับประคองจากภายนอก มาสู่การพึ่งพาตนเองและรับผิดชอบต่อชีวิตตนเองอย่างแท้จริง", bold_prefix="4) ลดการพึ่งพิงสิ่งแวดล้อมภายนอก (Environmental Support to Self-Support): ")
    add_bullet("บูรณาการส่วนของบุคลิกภาพที่เคยผลักไสเข้าเป็นองค์รวม ดำรงชีวิตอย่างเติมเต็มและงอกงาม", bold_prefix="5) ปรับปรุงความเป็นองค์รวมและความสมดุล (Wholeness & Integration): ")

    add_h2("7.2 บทบาทและหน้าที่หลักของผู้ให้คำปรึกษา")
    add_bullet("ทำหน้าที่เหมือนกระจกสะท้อน ชี้ชวนให้ผู้รับการปรึกษาตระหนักถึงสิ่งที่กำลังแสดงออก ทั้งน้ำเสียง ท่าทาง การสบตา และการหายใจ", bold_prefix="ส่งเสริมการตระหนักรู้: ")
    add_bullet("ไม่ยอมตกเป็นเครื่องมือในอุบายของผู้รับการปรึกษาที่ต้องการให้ผู้ให้คำปรึกษามารับผิดชอบแทน แต่จะสร้างความคับข้องใจอย่างเหมาะสมเพื่อให้ผู้รับคำปรึกษาเลิกพฤติกรรมพึ่งพิงและหันมาพึ่งตนเอง", bold_prefix="สร้างความคับข้องใจอย่างสร้างสรรค์ (Creative Frustration): ")
    add_bullet("เสนอแนะแบบฝึกหัดและการทดลองจริงในห้องให้คำปรึกษา (In-session Experiments) เพื่อให้ผู้รับคำปรึกษาได้เผชิญหน้ากับความรู้สึกใหม่ๆ ภายใต้บรรยากาศที่ปลอดภัย", bold_prefix="ออกแบบการทดลองเชิงประสบการณ์: ")
    add_bullet("ช่วยดึงศักยภาพและจุดเด่นของผู้รับการปรึกษาออกมาใช้ในการเผชิญสถานการณ์ใหม่ ควบคู่กับการสอนเทคนิคผ่อนคลาย เช่น การใส่ใจลมหายใจเข้า-ออก", bold_prefix="ดึงจุดแข็งและเสริมทักษะ: ")

    # Section 8
    add_h1("8. คลังเทคนิคและวิธีการให้คำปรึกษาแบบเกสตอลต์ (Gestalt Techniques & Experiments)")
    add_body("เทคนิคของเกสตอลต์เน้นการลงมือปฏิบัติจริง (Action-oriented) เพื่อกระตุ้นให้เกิดประสบการณ์ตรงและการหยั่งรู้ ได้แก่:")

    add_h2("8.1 การตระหนักรู้ต่อการสัมผัสและการกระทำ (Body & Sensory Awareness)")
    add_body(
        "ผู้ให้คำปรึกษาจะดึงความสนใจของผู้รับการปรึกษามาสู่สัญญาณทางกาย เช่น จังหวะการเต้นของหัวใจ ความตึงเกร็งที่ลำคอ อาการคัน หรือการหายใจตื้นๆ "
        "การสัมผัสทางกายที่เป็นกลางจะช่วยดึงบุคคลออกจากกรอบความคิดเดิมที่เคยถูกวางเงื่อนไขไว้ และนำไปสู่การกระทำที่เป็นธรรมชาติและสร้างสรรค์"
    )

    add_h2("8.2 การให้อยู่กับอารมณ์ความรู้สึก (Staying with the Feeling)")
    add_body(
        "เมื่อผู้รับการปรึกษาเริ่มมีอารมณ์เจ็บปวด เสียใจ หรือโกรธ ผู้ให้คำปรึกษาจะไม่รีบเปลี่ยนเรื่องหรือปลอบโยนเพื่อให้หยุดร้องไห้ "
        "แต่จะบอกให้ผู้รับการปรึกษา \"คงความรู้สึกนั้นไว้ และเฝ้าสังเกตว่าความรู้สึกนั้นกำลังนำพาเขาไปสู่เรื่องราวใด\" ซึ่งส่วนใหญ่จะนำไปสู่ปมคั่งค้างในอดีตที่แท้จริง"
    )

    add_h2("8.3 การตระหนักรู้ต่อความต้องการ (Needs Awareness)")
    add_body(
        "การถามตรงๆ ในจังหวะที่เหมาะสมว่า \"ในขณะนี้ คุณต้องการอะไร?\" คำถามนี้ช่วยดึงความต้องการที่จมอยู่ในพื้น (Ground) ขึ้นมาเป็นภาพ (Figure) "
        "ที่ชัดเจนในจิตสำนึก เพื่อกระตุ้นให้วางแผนลงมือกระทำอย่างรับผิดชอบ"
    )

    add_h2("8.4 การลงมือแสดงและพฤติกรรมที่ถูกกำกับ (Enactment & Directed Behavior)")
    add_body(
        "การเรียนรู้เกิดจากการลงมือทำจริง ผู้ให้คำปรึกษาจะให้ผู้รับการปรึกษาแสดงเหตุการณ์หรือพฤติกรรมที่ถูกสกัดกั้นไว้ "
        "เช่น กรณีผู้รับการปรึกษาชายที่คิดว่าตนเองต้องทำตัวสมบูรณ์แบบเพื่อให้แม่ยอมรับ ผู้ให้คำปรึกษาได้กำกับให้เขาเดินไปสบตาเพื่อนสมาชิกในกลุ่มทีละคนแล้วพูดว่า:"
    )
    add_box(
        "ตัวอย่างบทสนทนาพฤติกรรมที่ถูกกำกับ (Directed Behavior Dialogue)",
        "ผู้รับการปรึกษา: \"ผมรู้แล้วว่าจะต้องทำอย่างไร แต่ไม่น่าจำเป็นต้องพูดแบบนี้กับทุกคน\"\n"
        "ผู้ให้การปรึกษา: \"เอาน่า ลองดู พูดตามที่บอก\"\n"
        "ผู้รับการปรึกษา: (เริ่มสบตาด้วยความประหม่า) \"ผมไม่ต้องการทำอะไรเพื่อให้คุณชอบผม ผมดีพอเพราะผมเป็นตัวผมเอง...\"\n"
        "ผลที่เกิดขึ้น: เมื่อพูดกับสมาชิกหลายคน น้ำเสียงเริ่มมั่นคง ได้รับการตอบรับด้วยรอยยิ้มและการกอด จนเกิดการหยั่งรู้ว่า \"ตนเองมีคุณค่าได้โดยไม่ต้องพยายามเอาใจใคร\"",
        bg="F1F5F9", border="64748B", icon="💬"
    )

    add_h2("8.5 การใช้จินตนาการและเทคนิคเก้าอี้ว่างเปล่า (Empty Chair Technique)")
    add_body(
        "ใช้จัดการกับเรื่องราวคั่งค้างกับบุคคลที่ไม่สามารถมาพบได้ (เช่น พ่อแม่ที่ล่วงลับไปแล้ว หรือคนที่ตัดขาดความสัมพันธ์) "
        "โดยตั้งเก้าอี้ว่างไว้ตรงข้าม แล้วให้ผู้รับการปรึกษาจินตนาการว่าบุคคลนั้นนั่งอยู่ จากนั้นให้สื่อสาร 3 ขั้นตอนตามลำดับ:\n"
        "1. พูดถึงสิ่งที่ชอบหรือรู้สึกดีต่อบุคคลนั้นก่อน\n"
        "2. ปลดปล่อยความโกรธ ความขุ่นเคือง หรือความรู้สึกผิดที่คั่งค้างอยู่ออกมาให้หมดสิ้น\n"
        "3. กล่าวคำอำลาและปล่อยวาง ซึ่งเป็นสัญลักษณ์ของการปิดฉากความค้างคาในจิตใจ"
    )

    add_h2("8.6 การทำงานกับความฝันแบบเกสตอลต์ (Gestalt Dream Work)")
    add_body(
        "ต่างจากจิตวิเคราะห์ที่พยายามตีความหมายสัญลักษณ์ในระดับจิตไร้สำนึก เกสตอลต์มองว่า \"ความฝันคือการฉายภาพ (Projection) "
        "ของลักษณะนิสัยและบุคลิกภาพส่วนต่างๆ ของผู้ฝันเอง\" วิธีการบำบัดคือ ให้ผู้รับการปรึกษา **สมมุติตนเองเป็นตัวละครหรือสิ่งของทุกชิ้นในความฝัน** "
        "(เช่น หากฝันว่ากระต่ายวิ่งหนีสุนัขจิ้งจอกเข้าไปในโพรงไม้ ผู้รับการปรึกษาจะต้องสวมบทบาทเป็นทั้งกระต่าย สุนัขจิ้งจอก ทุ่งหญ้า และโพรงไม้) "
        "เพื่อสัมผัสอารมณ์ของทุกส่วน และบูรณาการลักษณะนิสัยเหล่านั้นกลับคืนสู่ตนเองอย่างเป็นเจ้าของ"
    )

    add_h2("8.7 การปะทะและบูรณาการระหว่าง Top Dog และ Underdog")
    add_body(
        "เกสตอลต์อธิบายความขัดแย้งภายในตนเองผ่านบุคลิกภาพ 2 ขั้วที่ทำสงครามกันอยู่เสมอ:",
        bold_prefix="ธรรมชาติของทั้งสองขั้ว: "
    )
    add_bullet("เปรียบเสมือน Superego หรือผู้ผดุงธรรม มีลักษณะเจ้าระเบียบ ชอบเทศนา สั่งสอน บงการว่า \"ควรทำอย่างนั้น ต้องทำอย่างนี้\" ดูมีอำนาจเข้มแข็ง", bold_prefix="Top Dog (จอมบงการ): ")
    add_bullet("ส่วนที่ไม่อยากทำ มีลักษณะแกล้งยอม ขอโทษขอโพย แก้ตัวว่าทำไม่ได้ แต่แท้จริงแล้ว \"ดื้อเงียบและไม่ยอมเปลี่ยน\" ซึ่งแท้จริง Underdog คือฝ่ายที่ควบคุมบุคคลอย่างแท้จริง", bold_prefix="Underdog (จอมตลบแตลง): ")
    add_body(
        "แนวทางแก้ไข: ให้ผู้รับการปรึกษานั่งสลับเก้าอี้เพื่อสวมบทบาทเป็น Top Dog และ Underdog สนทนาโต้เถียงกันเองอย่างเปิดเผย เพื่อให้ตระหนักรู้ถึงกลลวงของตนเอง "
        "และนำไปสู่การประนีประนอม บูรณาการทั้งสองส่วนเข้าด้วยกันอย่างสมดุล"
    )

    # Section 9: Key Takeaways Summary
    add_h1("9. สรุปสังเคราะห์จุดเน้นสำหรับทบทวนและเตรียมสอบ (Summary Cheat Sheet)")

    t_sum = doc.add_table(rows=7, cols=2)
    t_sum.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sum.autofit = False
    set_table_borders(t_sum, "CBD5E1", "4")
    s_w = [Inches(2.2), Inches(4.3)]
    for row in t_sum.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = s_w[idx]

    s_heads = ["หัวข้อหลัก", "สาระสำคัญที่ต้องจำแม่น"]
    for i, title in enumerate(s_heads):
        cell = t_sum.cell(0, i)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.name = "TH Sarabun New"
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(255, 255, 255)

    sum_data = [
        ("ความหมายของเกสตอลต์", "ส่วนรวมทั้งหมดที่มีความหมาย (A meaningful whole) กายและจิตเป็นองค์รวมแยกกันไม่ได้"),
        ("เวลาและสภาวะจิต", "ที่นี่และเดี๋ยวนี้ (Here & Now) การตระหนักรู้ (Awareness) มีพลังเยียวยาในตนเอง"),
        ("ภาพและพื้น (Figure & Ground)", "ภาพ = ความต้องการเด่นชัด / พื้น = สิ่งแวดล้อม สุขภาพจิตดีต้องสลับไปมาอย่างยืดหยุ่น"),
        ("4 กลไกขัดขวางการติดต่อ", "1. กลืนค่านิยม (ไม่เคี้ยว) | 2. ฉายภาพ (โยนให้คนอื่น) | 3. ย้อนสู่ตน (ทำร้ายตัวเอง) | 4. หละหลวม (กลืนเป็นเนื้อเดียว)"),
        ("5 ชั้นของอาการประสาท", "เสแสร้ง ➔ กลัว ➔ ทางตัน ➔ ระเบิดทางตัน ➔ ระเบิดสู่ภายนอก (มีชีวิตชีวา/องค์รวม)"),
        ("เทคนิคการบำบัดเด่น", "เก้าอี้ว่างเปล่า (Empty Chair), อยู่กับความรู้สึก (Staying with feeling), สวมบทบาทในฝัน (Dream work), บูรณาการ Top dog vs Underdog"),
    ]

    for row_idx, (col0, col1) in enumerate(sum_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        c0 = t_sum.cell(row_idx, 0)
        c1 = t_sum.cell(row_idx, 1)
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_margins(c0, top=100, bottom=100, left=140, right=140)
        set_cell_margins(c1, top=100, bottom=100, left=140, right=140)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(col0)
        r0.bold = True
        r0.font.name = "TH Sarabun New"
        r0.font.size = Pt(13.5)
        r0.font.color.rgb = RGBColor(30, 58, 138)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(col1)
        r1.font.name = "TH Sarabun New"
        r1.font.size = Pt(13.5)
        r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Save documents
    out_name1 = "สรุปบทที่11_ทฤษฎีการให้คำปรึกษาแบบเกสตอลต์.docx"
    doc.save(out_name1)
    print(f"Saved: {out_name1}")

    # Also copy to งาน/ folder
    target_dir = "งาน"
    if os.path.exists(target_dir):
        out_name2 = os.path.join(target_dir, "สรุปบทที่11_ทฤษฎีการให้คำปรึกษาแบบเกสตอลต์.docx")
        shutil.copyfile(out_name1, out_name2)
        print(f"Copied to: {out_name2}")

if __name__ == '__main__':
    create_summary_doc()
