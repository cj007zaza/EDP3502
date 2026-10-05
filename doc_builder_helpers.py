import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_callout_border(cell, border_hex="2563EB", border_size="36"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def add_callout(doc, quote_title, quote_text, plain_explanation, action_text="", connection_text=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)

    cell = table.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_callout_border(cell, "2563EB", "32")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.line_spacing = 1.15
    run_quote_title = p1.add_run(f"📜 ข้อความต้นฉบับ / นิยามวิชาการ: {quote_title}\n")
    run_quote_title.bold = True
    run_quote_title.font.name = "TH Sarabun New"
    run_quote_title.font.size = Pt(14)
    run_quote_title.font.color.rgb = RGBColor(30, 58, 138)

    run_quote = p1.add_run(f"\"{quote_text}\"\n")
    run_quote.italic = True
    run_quote.font.name = "TH Sarabun New"
    run_quote.font.size = Pt(13)
    run_quote.font.color.rgb = RGBColor(71, 85, 105)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(6)
    p2.paragraph_format.space_after = Pt(4)
    p2.paragraph_format.line_spacing = 1.15
    run_exp_title = p2.add_run("💡 ถอดรหัสคำอธิบายภาษาคน & ภาพจำ:\n")
    run_exp_title.bold = True
    run_exp_title.font.name = "TH Sarabun New"
    run_exp_title.font.size = Pt(14)
    run_exp_title.font.color.rgb = RGBColor(180, 83, 9)

    run_exp = p2.add_run(plain_explanation + "\n")
    run_exp.font.name = "TH Sarabun New"
    run_exp.font.size = Pt(13.5)
    run_exp.font.color.rgb = RGBColor(30, 41, 59)

    if action_text:
        p3 = cell.add_paragraph()
        p3.paragraph_format.space_before = Pt(4)
        p3.paragraph_format.space_after = Pt(2)
        p3.paragraph_format.line_spacing = 1.15
        run_act_title = p3.add_run("🛠️ ครูนำไปใช้ทำอะไรหน้างาน: ")
        run_act_title.bold = True
        run_act_title.font.name = "TH Sarabun New"
        run_act_title.font.size = Pt(13)
        run_act_title.font.color.rgb = RGBColor(13, 148, 136)

        run_act = p3.add_run(action_text)
        run_act.font.name = "TH Sarabun New"
        run_act.font.size = Pt(13)
        run_act.font.color.rgb = RGBColor(51, 65, 85)

    if connection_text:
        p4 = cell.add_paragraph()
        p4.paragraph_format.space_before = Pt(4)
        p4.paragraph_format.space_after = Pt(2)
        p4.paragraph_format.line_spacing = 1.15
        run_conn_title = p4.add_run("🔗 เชื่อมโยงกับ: ")
        run_conn_title.bold = True
        run_conn_title.font.name = "TH Sarabun New"
        run_conn_title.font.size = Pt(13)
        run_conn_title.font.color.rgb = RGBColor(124, 58, 237)

        run_conn = p4.add_run(connection_text)
        run_conn.font.name = "TH Sarabun New"
        run_conn.font.size = Pt(13)
        run_conn.font.color.rgb = RGBColor(51, 65, 85)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(6)

print("Helper functions ready.")
