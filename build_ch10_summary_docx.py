import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_ch10_summary_document(output_path):
    doc = docx.Document()

    # Page setup - Margins 1 inch (72 pt / 1440 dxa)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.27)  # A4 standard
        section.page_height = Inches(11.69)

    # Base font setting
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'TH Sarabun New'
    normal_style.font.size = Pt(15)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)  # Slate 800
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # XML Helper functions
    def set_font_run(run, name='TH Sarabun New', size_pt=15, bold=False, italic=False, color_rgb=None):
        run.font.name = name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        if color_rgb:
            run.font.color.rgb = color_rgb
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{name}" w:hAnsi="{name}" w:cs="{name}"/>')
        rPr.append(rFonts)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                         color="CBD5E1", sz="8"):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

    # Document Header / Banner
    banner_table = doc.add_table(rows=1, cols=1)
    banner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner_table.autofit = False
    banner_table.columns[0].width = Inches(6.27)
    
    b_cell = banner_table.cell(0, 0)
    set_cell_background(b_cell, "1E3A8A")  # Navy Blue
    set_cell_margins(b_cell, top=240, bottom=240, left=240, right=240)
    set_cell_borders(b_cell)

    bp = b_cell.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bp.paragraph_format.space_after = Pt(2)
    r_sub = bp.add_run("เอกสารสรุปเนื้อหารายวิชา EDP3502 การให้คำปรึกษาและการช่วยเหลือผู้เรียน\n")
    set_font_run(r_sub, size_pt=14, bold=False, color_rgb=RGBColor(219, 234, 254))

    r_title = bp.add_run("สรุปบทที่ 10: การให้คำปรึกษาแบบบุคคลเป็นศูนย์กลาง\n")
    set_font_run(r_title, size_pt=22, bold=True, color_rgb=RGBColor(255, 255, 255))

    r_eng = bp.add_run("Person-Centered Counseling / Therapy (Carl Rogers)")
    set_font_run(r_eng, size_pt=14, italic=True, color_rgb=RGBColor(191, 219, 254))

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(8)
    p_spacer.paragraph_format.space_after = Pt(4)

    # Helper text functions
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        set_font_run(run, size_pt=18, bold=True, color_rgb=RGBColor(30, 58, 138))
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        set_font_run(run, size_pt=16, bold=True, color_rgb=RGBColor(37, 99, 235))
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        set_font_run(run, size_pt=15, bold=True, color_rgb=RGBColor(15, 118, 110))
        return p

    def add_body(text, bold_prefix="", indent=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if indent > 0:
            p.paragraph_format.left_indent = Inches(0.25 * indent)
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            set_font_run(run_b, size_pt=15, bold=True, color_rgb=RGBColor(15, 23, 42))
        run_t = p.add_run(text)
        set_font_run(run_t, size_pt=15, bold=False, color_rgb=RGBColor(51, 65, 85))
        return p

    def add_bullet(text, bold_prefix="", level=0):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.25 + 0.2 * level)
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            set_font_run(run_b, size_pt=15, bold=True, color_rgb=RGBColor(15, 23, 42))
        run_t = p.add_run(text)
        set_font_run(run_t, size_pt=15, bold=False, color_rgb=RGBColor(51, 65, 85))
        return p

    def add_callout_box(title, text, icon="💡", bg_color="F8FAFC", border_color="2563EB"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(6.27)
        cell = table.cell(0, 0)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        set_cell_borders(cell, left="single", color=border_color, sz="36")
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run_title = p.add_run(f"{icon} {title}\n")
        set_font_run(run_title, size_pt=15, bold=True, color_rgb=RGBColor(30, 58, 138))
        
        run_body = p.add_run(text)
        set_font_run(run_body, size_pt=14.5, bold=False, color_rgb=RGBColor(30, 41, 59))
        
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(2)
        p_sp.paragraph_format.space_after = Pt(4)

    def add_dialogue_box(speaker_roles, dialogue_lines, analysis_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(6.27)
        cell = table.cell(0, 0)
        set_cell_background(cell, "F1F5F9")  # Slate 100
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        set_cell_borders(cell, left="single", color="0D9488", sz="32")  # Teal border
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        run_head = p.add_run("💬 ตัวอย่างบทสนทนา & การวิเคราะห์เทคนิค\n")
        set_font_run(run_head, size_pt=14.5, bold=True, color_rgb=RGBColor(13, 148, 136))
        
        for spk, line in dialogue_lines:
            p_line = cell.add_paragraph()
            p_line.paragraph_format.space_before = Pt(1)
            p_line.paragraph_format.space_after = Pt(2)
            p_line.paragraph_format.line_spacing = 1.1
            run_spk = p_line.add_run(f"{spk}: ")
            set_font_run(run_spk, size_pt=14, bold=True, color_rgb=RGBColor(30, 41, 59))
            run_c = p_line.add_run(f"\"{line}\"")
            set_font_run(run_c, size_pt=14, italic=True, color_rgb=RGBColor(71, 85, 105))
            
        p_an = cell.add_paragraph()
        p_an.paragraph_format.space_before = Pt(6)
        p_an.paragraph_format.space_after = Pt(2)
        run_at = p_an.add_run("🔍 การวิเคราะห์ทางจิตวิทยา: ")
        set_font_run(run_at, size_pt=14, bold=True, color_rgb=RGBColor(180, 83, 9))
        run_ac = p_an.add_run(analysis_text)
        set_font_run(run_ac, size_pt=14, color_rgb=RGBColor(30, 41, 59))

        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(2)
        p_sp.paragraph_format.space_after = Pt(4)

    # ------------------ SECTION 1 ------------------
    add_h1("1. ประวัติและแนวคิดพื้นฐานเกี่ยวกับมนุษย์")
    add_body("ผู้ก่อตั้งทฤษฎีคือ คาร์ล โรเจอรส์ (Carl Rogers) นักจิตวิทยาชาวอเมริกันผู้มีประสบการณ์ทำการบำบัดรักษาผู้รับการปรึกษามากกว่า 30 ปี (เฉลี่ยสัปดาห์ละ 15–20 ชั่วโมง) โดยมีมุมมองเชิงมนุษยนิยม (Humanistic Perspective) ต่อศักยภาพของมนุษย์ ดังนี้:")
    
    add_bullet("ทุกคนมีพลังขับเคลื่อนตามธรรมชาติที่จะพัฒนาตนเองไปสู่ความงอกงามสูงสุด มีความสามารถในการเรียนรู้ แก้ปัญหา และปรับปรุงตนเองให้ดีขึ้นได้เสมอ", bold_prefix="ศักยภาพในตนเอง (Self-Actualizing Tendency): ")
    add_bullet("ผู้ที่ไม่เชื่อมั่นในตนเอง ไม่ได้เกิดจากความบกพร่องแต่กำเนิด แต่เกิดจากประสบการณ์ในอดีตที่ถูกบิดเบือน รูปแบบการอบรมเลี้ยงดูในครอบครัว ความคาดหวังจากโรงเรียน และแรงกดดันทางสังคม", bold_prefix="สาเหตุของการสูญเสียความเชื่อมั่น: ")
    add_bullet("ตราบใดที่บุคคลยินดีที่จะ 'ยอมรับว่าตนเองต้องเป็นผู้รับผิดชอบชีวิตของตน' บุคคลจะเริ่มสัมผัสได้ถึงอิสรภาพเหมือนเกิดใหม่ กลายเป็นผู้กำหนดชะตาชีวิตของตนเองได้อย่างแจ่มชัด", bold_prefix="กุญแจสำคัญสู่การหลุดพ้น: ")

    add_callout_box(
        "หัวใจสำคัญของ Carl Rogers",
        "โรเจอรส์เชื่อว่าผู้รับการปรึกษาไม่ใช่ 'คนป่วย' หรือ 'คนไร้ความสามารถ' ที่ต้องรอรับคำสั่งสอน แต่เป็นบุคคลผู้มีทรัพยากรภายในครบถ้วน หน้าที่ของผู้ให้คำปรึกษาคือสร้างบรรยากาศที่เอื้อต่อการงอกงาม เพื่อให้ศักยภาพนั้นได้เผยตัวออกมาเอง",
        icon="🌱", bg_color="EFF6FF", border_color="2563EB"
    )

    # ------------------ SECTION 2 ------------------
    add_h1("2. ทฤษฎีบุคลิกภาพ (Personality Theory)")
    add_body("ทฤษฎีบุคลิกภาพของโรเจอรส์ประกอบด้วยโครงสร้างและแนวคิดสำคัญ 2 องค์ประกอบ ได้แก่ อินทรีย์ (Organism) และ ตน (Self):")

    add_h2("2.1 อินทรีย์ (Organism)")
    add_bullet("คือผลรวมของประสบการณ์ทั้งหมดที่อินทรีย์เผชิญ ซึ่งเฉพาะตัวบุคคลผู้นั้นเท่านั้นที่จะเข้าใจได้ ประสบการณ์นี้มีทั้งส่วนที่รู้ตัว (จิตสำนึก) และไม่รู้ตัว (จิตไร้สำนึก)", bold_prefix="สนามของปรากฏการณ์ (Phenomenological Field): ")
    add_bullet("มนุษย์ไม่ได้แสดงพฤติกรรมตามความจริงเชิงวัตถุภายนอก แต่แสดงออกตาม 'ความจริงที่อิงตนเองเป็นหลัก' ซึ่งอาจตรงหรือไม่ตรงกับความจริงภายนอกก็ได้", bold_prefix="ความจริงที่เป็นอัตวิสัย (Subjective Reality): ")
    add_bullet("หากบุคคลตรวจสอบความจริงอัตวิสัยอย่างไม่เหมาะสม หรือปักใจเชื่อมั่นอย่างรวดเร็วเกินไปโดยไม่พิจารณาว่าเป็นเพียงสมมุติฐาน ผลคือจะเกิดการรับรู้เกี่ยวกับตนเองอย่างผิดพลาด", bold_prefix="ข้อพึงระวัง: ")

    add_h2("2.2 ตน (Self / Self-concept)")
    add_body("คือส่วนที่พัฒนาแยกออกมาจากอินทรีย์ เป็นภาพลักษณ์และโครงสร้างความคิดว่า 'ฉันคือใคร' หรือ 'ข้าพเจ้าเป็นคนอย่างไร'")
    add_bullet("กระบวนการแปลงประสบการณ์ให้กลายเป็นการรับรู้ หากยอมรับความจริงจะให้สัญลักษณ์ตรงไปตรงมา แต่ถ้าประสบการณ์ขัดแย้งกับภาพลักษณ์ตน บุคคลจะบิดเบือนหรือปฏิเสธความรู้สึกที่แท้จริง สิ่งที่ถูกปฏิเสธจะถูกผลักลงไปสะสมไว้ในจิตไร้สำนึก", bold_prefix="การให้สัญลักษณ์กับประสบการณ์ (Symbolization): ")
    add_bullet("เกิดขึ้นเมื่อภาพลักษณ์ของตน (Self) กับประสบการณ์จริงของร่างกายและอารมณ์ (Organismic Experience) ไม่สอดคล้องกัน บุคคลจะรู้สึกถูกคุกคาม จิตไร้สำนึกจะรับรู้ก่อนจิตสำนึก และส่งสัญญาณเตือนออกมาทางสรีระ", bold_prefix="ความไม่สอดคล้อง (Incongruence): ")

    # Table of Incongruence Symptoms
    t_inc = doc.add_table(rows=4, cols=2)
    t_inc.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_inc.autofit = False
    t_inc.columns[0].width = Inches(2.2)
    t_inc.columns[1].width = Inches(4.07)
    
    headers = ["มิติที่ได้รับผลกระทบ", "อาการและพฤติกรรมที่แสดงออก"]
    for i, h in enumerate(headers):
        c = t_inc.cell(0, i)
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, 100, 100, 140, 140)
        set_cell_borders(c, color="CBD5E1")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_font_run(r, size_pt=14, bold=True, color_rgb=RGBColor(255, 255, 255))
        
    inc_data = [
        ("1. ผลกระทบทางสรีระ (Physiological)", "ร่างกายตึงเครียด, หัวใจเต้นเร็ว, มือสั่น, หลบสายตา, เหงื่อออก, ปวดศีรษะเรื้อรัง"),
        ("2. ผลกระทบทางอารมณ์ (Emotional)", "รู้สึกกระวนกระวาย, เบื่อหน่าย, หงุดหงิดง่าย, หวั่นกลัว, วิตกกังวลลึกๆ"),
        ("3. ผลกระทบทางความคิด (Cognitive)", "เกิดความสับสน, ไม่รู้ว่าแท้จริงตนเองต้องการอะไร, รู้สึกขัดแย้งในตนเองตลอดเวลา")
    ]
    for row_idx, (dim, desc) in enumerate(inc_data, start=1):
        c0 = t_inc.cell(row_idx, 0)
        c1 = t_inc.cell(row_idx, 1)
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_margins(c0, 100, 100, 140, 140)
        set_cell_margins(c1, 100, 100, 140, 140)
        set_cell_borders(c0, color="E2E8F0")
        set_cell_borders(c1, color="E2E8F0")
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(dim)
        set_font_run(r0, size_pt=14, bold=True, color_rgb=RGBColor(30, 41, 59))
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(desc)
        set_font_run(r1, size_pt=14, color_rgb=RGBColor(71, 85, 105))

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(4)
    p_sp.paragraph_format.space_after = Pt(4)

    add_h2("2.3 การรวบรวมภาพลักษณ์ของตนใหม่ (Reorganizing Self)")
    add_body("คือกระบวนการฟื้นคืนความสอดคล้องระหว่างตนกับประสบการณ์จริง โดยมีสาระสำคัญดังนี้:")
    add_bullet("บุคคลต้องค่อยๆ ยอมรับความรู้สึกและประสบการณ์ที่เคยปฏิเสธหรือเก็บกดไว้ในจิตไร้สำนึก ซึ่งต้องอาศัยบรรยากาศที่ 'เป็นอิสระและปราศจากการถูกประเมินตัดสิน' (Freedom from evaluation)", bold_prefix="การก้าวผ่านหน้ากากที่สวมอยู่: ")
    add_bullet("บุคคลค้นพบและยอมรับความจริงว่า ตนเองมีทั้งด้านบวกและด้านลบ เป็นทั้ง 'เทวดาและปีศาจ' มีทั้งความโกรธ เกลียด อิจฉา ชิงดีชิงเด่น ควบคู่กับความรัก อบอุ่น เสียสละ และเมตตา", bold_prefix="เกณฑ์ความสำเร็จในการค้นพบตน: ")
    add_bullet("เมื่อบุคคลยอมรับความจริงของตนได้อย่างครบถ้วน อาการทางสรีระผิดปกติจะยุติลง และไม่ได้นำไปสู่พฤติกรรมก้าวร้าวทำลายล้าง แต่กลับเป็นจุดเริ่มต้นของการพัฒนาสู่พฤติกรรมสร้างสรรค์เชิงบวก", bold_prefix="ผลลัพธ์เชิงบวก: ")

    # ------------------ SECTION 3 ------------------
    add_h1("3. กระบวนการสำคัญ 4 ประการในการพัฒนาการของบุคลิกภาพ")
    add_body("ทิศทางการพัฒนาบุคลิกภาพของบุคคลขึ้นอยู่กับปฏิสัมพันธ์ของ 4 กระบวนการหลัก ดังตารางสรุป:")

    t_proc = doc.add_table(rows=5, cols=3)
    t_proc.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_proc.autofit = False
    t_proc.columns[0].width = Inches(1.8)
    t_proc.columns[1].width = Inches(2.6)
    t_proc.columns[2].width = Inches(1.87)

    proc_headers = ["กระบวนการพัฒนา", "สาระสำคัญและกลไกการทำงาน", "ผลกระทบต่อบุคคล"]
    for i, h in enumerate(proc_headers):
        c = t_proc.cell(0, i)
        set_cell_background(c, "0F766E")  # Teal Header
        set_cell_margins(c, 100, 100, 120, 120)
        set_cell_borders(c, color="CBD5E1")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_font_run(r, size_pt=14, bold=True, color_rgb=RGBColor(255, 255, 255))

    proc_data = [
        ("1. กระบวนการให้คุณค่าของอินทรีย์\n(Organismic Valuing Process)", 
         "การประเมินสิ่งต่างๆ จากผัสสะและความรู้สึกภายในของตนเองโดยตรงตามธรรมชาติ ไม่เสแสร้ง (เหมือนทารกที่หิวกิน อิ่มหยุด ไม่ยึดเกณฑ์สังคมภายนอก)", 
         "นำไปสู่การพัฒนาตนเองอย่างเป็นอิสระและใช้ศักยภาพสูงสุด"),
        ("2. ความต้องการการยอมรับจากผู้อื่น\n(Need for Positive Regard)", 
         "เด็กต้องพึ่งพาผู้ใหญ่เพื่อความอยู่รอด จึงเริ่มสังเกตว่าพฤติกรรมใดผู้ใหญ่ชม พฤติกรรมใดผู้ใหญ่ไม่ชอบ และเริ่มปรับตัวเอาใจผู้อื่น", 
         "เริ่มสงสัย ไม่มั่นใจในความรู้สึกตนเอง และยึดผู้อื่นเป็นหลัก"),
        ("3. ความต้องการการยอมรับตนเอง\n(Need for Self-Regard)", 
         "การรับรู้คุณค่าของตนเองพัฒนามาจากการได้รับการยอมรับจากผู้อื่น หากผู้ใหญ่ไม่ยอมรับ เด็กจะไม่ยอมรับตนเองด้วย", 
         "ทำให้กระบวนการให้คุณค่าของอินทรีย์เริ่มแผ่วกำลังลง"),
        ("4. เงื่อนไขแห่งความมีคุณค่า\n(Conditions of Worth)", 
         "บุคคลตั้งกฎเกณฑ์ว่า 'ต้องทำตามความคาดหวังของสังคมจึงจะมีคุณค่า' ปฏิเสธความต้องการแท้จริงของตนเพื่อแลกกับการยอมรับ", 
         "ขัดขวางการทำงานของศักยภาพอินทรีย์ นำไปสู่บุคลิกภาพที่แปลกแยก")
    ]
    for row_idx, (col0, col1, col2) in enumerate(proc_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for c_idx, text_val in enumerate([col0, col1, col2]):
            cell = t_proc.cell(row_idx, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 120, 120)
            set_cell_borders(cell, color="E2E8F0")
            p = cell.paragraphs[0]
            r = p.add_run(text_val)
            bold_flag = True if c_idx == 0 else False
            set_font_run(r, size_pt=13.5, bold=bold_flag, color_rgb=RGBColor(30, 41, 59))

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(4)
    p_sp.paragraph_format.space_after = Pt(4)

    # ------------------ SECTION 4 & 5 ------------------
    add_h1("4. เป้าหมายของการให้คำปรึกษา")
    add_bullet("สร้างสถานการณ์และเงื่อนไขที่ปลอดภัย ไม่ตัดสิน เพื่อให้ผู้รับการปรึกษากล้าเปิดเผยตนเองและสำรวจความรู้สึกที่แท้จริง", bold_prefix="1) สร้างสภาวะที่ปราศจากการประเมิน (Non-evaluative climate): ")
    add_bullet("ส่งเสริมให้บุคคลทบทวนตนเอง ยอมรับประสบการณ์ทั้งหมด ทำให้ตน (Self) กับประสบการณ์ (Experience) สอดคล้องตรงกัน กลายเป็น 'บุคคลใหม่ที่สมบูรณ์'", bold_prefix="2) ขจัดความไม่สอดคล้องในตนเอง (Resolving Incongruence): ")
    add_bullet("ช่วยให้ผู้รับการปรึกษามีสติสัมผัสรับรู้อารมณ์ความรู้สึกที่เกิดขึ้นในปัจจุบันขณะ ผ่านการสะท้อนความรู้สึกอย่างลึกซึ้ง", bold_prefix="3) จดจ่ออยู่กับประสบการณ์ปัจจุบัน (Here and Now): ")

    add_h1("5. บทบาทและหน้าที่ของผู้ให้คำปรึกษา")
    add_bullet("ทำความเข้าใจโลกและมุมมองของผู้รับการปรึกษาโดยปราศจากสมมุติฐานหรืออคติล่วงหน้า")
    add_bullet("สร้างสัมพันธภาพที่มีความไว้วางใจซึ่งกันและกัน ยอมรับ และมีความเป็นธรรมชาติ")
    add_bullet("ไม่ใส่ ไม่ป้อน หรือยัดเยียดค่านิยม มาตรฐาน หรือคำสั่งสอนใดๆ แก่ผู้รับการปรึกษา")
    add_bullet("รับฟังและตอบสนองอย่างเอาใจใส่ สื่อสารความเคารพและการเห็นคุณค่าทั้งทางวาจาและอวัจนภาษา")
    add_bullet("แสดงความไว้วางใจและเชื่อมั่นในศักยภาพ การกำกับดูแล และความรับผิดชอบในชีวิตของตัวผู้รับการปรึกษาเอง")
    add_bullet("ส่งเสริมบรรยากาศที่เปิดกว้างและเอื้อให้ผู้รับการปรึกษากล้าทดลองแสดงพฤติกรรมใหม่ๆ")

    # ------------------ SECTION 6 ------------------
    add_h1("6. เงื่อนไขหลัก 3 ประการและเทคนิคการให้คำปรึกษา")
    add_body("โรเจอรส์ไม่เน้นกลวิธีหรือแบบแผนทางเทคนิคที่ตายตัว แต่เน้น 'คุณภาพของสัมพันธภาพ' ผ่านเงื่อนไขจำเป็น 3 ประการ:")

    add_h2("6.1 ความเข้าใจอย่างลึกซึ้ง (Empathy)")
    add_body("คือความสามารถในการก้าวเข้าสู่โลกภายในของผู้รับการปรึกษา มองเห็นและรู้สึกเสมือนเป็นตัวเขาเอง โดยไม่สูญเสียความเป็นตัวของผู้ให้คำปรึกษา ประกอบด้วยทักษะสำคัญ:")
    add_bullet("พักวางเรื่องราวส่วนตัว สบสายตา พยักหน้า แสดงสีหน้าที่ตอบรับกับอารมณ์ ระวังระยะห่างการนั่ง และใช้น้ำเสียงที่อบอุ่น ไม่ขุ่นมัว ไม่กดทับอำนาจ (น้ำเสียงมีความสำคัญยิ่งต่อผู้รับการปรึกษาที่พิการทางสายตา)", bold_prefix="การตั้งใจจดจ่อ (Attending): ")
    add_bullet("การสื่อสารสะท้อนกลับทำได้ 2 ระดับ: (1) การเรียบเรียงถ้อยคำใหม่ (Paraphrasing) ซึ่งเป็นการจับความระดับผิวเผิน และ (2) การสะท้อนความรู้สึก (Reflecting) ซึ่งมีพลังมากกว่า เพราะช่วยให้เห็นความคาดหวังในตนเอง ไม่ติดอยู่กับบทบาทเหยื่อผู้ถูกกระทำ", bold_prefix="การสื่อสารด้วยคำพูด: ")
    add_bullet("สังเกตสีหน้า แววตา ริมฝีปาก การใช้มือเท้า รวมถึง 'เนื้อหาที่ผู้รับการปรึกษาจงใจข้ามหรือหลีกเลี่ยงไม่พูดถึง' และการเลือกตำแหน่งที่นั่ง", bold_prefix="การสังเกตอวัจนภาษา: ")
    add_bullet("ความเงียบไม่ใช่ความอึดอัด แต่เป็นช่วงเวลาสำคัญที่ผู้รับการปรึกษากำลังจมอยู่ในภวังค์ความคิดความรู้สึก ผู้ให้คำปรึกษาต้องรอคอยด้วยความเคารพ ไม่เร่งรีบตัดบท เพื่อส่งสัญญาณว่าไม่ต้องการควบคุมเขา", bold_prefix="การใช้ความเงียบ (Silence): ")

    add_h2("6.2 ความจริงแท้ / ความสอดคล้องกลมกลืน (Genuineness / Congruence)")
    add_body("คือการแสดงออกที่ตรงกับความรู้สึกภายใน ซื่อตรง ไม่เสแสร้ง ประกอบด้วย 5 คุณลักษณะ:")
    add_bullet("ไม่สวมหน้ากากนักวิชาการหรือผู้เชี่ยวชาญ สื่อความคิดความรู้สึกต่อสิ่งที่เกิดขึ้นอย่างตรงไปตรงมา", bold_prefix="1) เป็นอิสระจากบทบาท (Role-free): ")
    add_bullet("สื่อสารอย่างมั่นใจ ผ่อนคลาย ไม่ต้องเกร็งหรือชั่งน้ำหนักคำพูดในใจตลอดเวลา", bold_prefix="2) มีความเป็นธรรมชาติ (Spontaneous): ")
    add_bullet("ตระหนักรู้ข้อดีและข้อจำกัดของตนเอง พร้อมเปิดรับอารมณ์ลบหรือคำตำหนิโดยไม่แก้ตัวหรือโต้กลับ", bold_prefix="3) ปราศจากการปกป้องตนเอง (Non-defensive): ")
    add_bullet("เสมอต้นเสมอปลาย คำพูดตรงกับการกระทำ ไม่แสดงพฤติกรรมที่ขัดกับค่านิยมตนเอง (ไม่สองหน้า)", bold_prefix="4) สม่ำเสมอคงเส้นคงวา (Consistent): ")
    add_bullet("แบ่งปันประสบการณ์หรือความรู้สึกของตนเองเมื่อจังหวะเวลาเหมาะสมและเป็นประโยชน์ต่อผู้รับคำปรึกษา", bold_prefix="5) การแบ่งปันตน (Self-disclosure): ")

    add_h2("6.3 การยอมรับทางบวกอย่างไม่มีเงื่อนไข (Unconditional Positive Regard)")
    add_body("คือการยอมรับนับถือในคุณค่าและศักดิ์ศรีความเป็นมนุษย์ของผู้รับการปรึกษาอย่างบริสุทธิ์ใจ โดยปราศจากการตัดสินประเมิน:")
    add_bullet("เอาใจใส่ต่อเรื่องราวและความรู้สึกอย่างจริงจัง", bold_prefix="พฤติกรรมที่ 1: ")
    add_bullet("สื่อทัศนคติที่ปราศจากการตัดสินหรือประเมินค่า (Non-judgmental attitude)", bold_prefix="พฤติกรรมที่ 2: ")
    add_bullet("ตอบสนองด้วยความเข้าใจเสมือนอยู่ในกรอบอ้างอิงเดียวกัน (Internal frame of reference)", bold_prefix="พฤติกรรมที่ 3: ")
    add_bullet("แสดงให้เห็นว่าผู้รับการปรึกษามีศักยภาพและความสามารถที่จะตัดสินใจและลงมือแก้ไขชีวิตตนเอง", bold_prefix="พฤติกรรมที่ 4: ")

    # ------------------ SECTION 7 ------------------
    add_h1("7. ตัวอย่างบทสนทนาและการวิเคราะห์เทคนิคการให้คำปรึกษา")
    
    # Dialogue 1
    add_h2("ตัวอย่างที่ 1: การเรียบเรียงถ้อยคำ (Paraphrasing) vs การสะท้อนความรู้สึก (Reflecting)")
    add_dialogue_box(
        speaker_roles=["ผู้รับฯ", "ผู้ให้ฯ"],
        dialogue_lines=[
            ("ผู้รับการปรึกษา", "หนูรู้สึกค่อนข้างแย่หลังจากสอบ ก่อนหน้านั้นไม่คิดว่าจะทำไม่ได้"),
            ("ผู้ให้ฯ (เรียบเรียง)", "หนูรู้สึกผิดหวังกับการสอบ"),
            ("ผู้ให้ฯ (สะท้อนลึก)", "หนูรู้สึกนึกไม่ถึงและผิดหวังในการทำข้อสอบ เพราะหนูคาดคิดมาก่อนว่าจะทำได้")
        ],
        analysis_text="การเรียบเรียงถ้อยคำทำให้ผู้รับบริการรับรู้เพียงความผิดหวังในสถานะ 'เหยื่อผู้ถูกกระทำ' แต่การสะท้อนความรู้สึกในระดับก้าวหน้า ดึงประเด็น 'ความคาดหวังในตนเอง' ขึ้นมาให้ตระหนักรู้ ซึ่งเป็นสิ่งที่ผู้รับบริการสามารถควบคุมและจัดการได้ดีกว่า"
    )

    # Dialogue 2
    add_h2("ตัวอย่างที่ 2: การเน้นความกังวลในปัจจุบัน vs การหลงไปกับบุคคลภายนอก")
    add_dialogue_box(
        speaker_roles=["ผู้รับฯ", "ผู้ให้ฯ"],
        dialogue_lines=[
            ("ผู้รับการปรึกษา", "พ่อกับแม่ไม่มีความไว้วางใจหนูเลย หนูไม่เคยทำอะไรถูกต้องเลยในสายตาของท่าน"),
            ("ผู้ให้ฯ (หลงประเด็น)", "พ่อแม่หนูไม่มีความมั่นใจในตัวหนู พ่อแม่หลายคนก็เป็นเช่นนั้นกับลูกๆ"),
            ("ผู้ให้ฯ (เน้นปัจจุบัน)", "หนูรู้สึกเสียใจและโกรธเคืองพ่อแม่ เพราะหนูอยากจะให้พ่อแม่แสดงความมั่นใจในตัวหนูบ้าง หนูอยากจะทำอะไรที่เป็นการพิสูจน์ว่าหนูสามารถทำอะไรได้ถูกต้อง")
        ],
        analysis_text="ผู้ให้คำปรึกษาดึงประเด็นกลับมาที่ 'ความรู้สึกและความต้องการของผู้รับบริการในปัจจุบัน' แทนที่จะเปลี่ยนเรื่องไปคุยถึงพ่อแม่ซึ่งไม่ได้อยู่ในห้องปรึกษา เป็นการโอนความรับผิดชอบในการสำรวจตนเองกลับมาที่ตัวผู้รับการปรึกษาโดยตรง"
    )

    # Dialogue 3
    add_h2("ตัวอย่างที่ 3: การสื่อความจริงแท้เมื่อเผชิญหน้ากับความขัดแย้ง (Genuineness)")
    add_dialogue_box(
        speaker_roles=["ผู้รับฯ", "ผู้ให้ฯ"],
        dialogue_lines=[
            ("ผู้รับการปรึกษา", "อาจารย์พูดตัดบท ไม่ยอมเป็นประธานวิทยานิพนธ์ให้ผม แล้วมาสอนในห้องว่าต้องแคร์ผู้รับการปรึกษา อาจารย์ทำตัวเหมือนคนตีสองหน้า!"),
            ("ผู้ให้การปรึกษา", "ครูมองออกว่าเธอยังคงโกรธครูอยู่ เพราะการกระทำของครูไม่ตรงกับที่เธอคาดหวัง... ครูเสียใจที่เธอรู้สึกอย่างนั้น และถ้าครูจะยอมรับเป็นประธานฯ ของเธอในตอนนี้ ครูเองก็คงจะรู้สึกว่าตนเองกลายเป็นคนตีสองหน้าจริงๆ และรวมทั้งตัวเธอด้วย")
        ],
        analysis_text="ผู้ให้คำปรึกษาไม่แก้ตัว ไม่ปกป้องตนเอง (Non-defensive) ไม่ปฏิเสธอารมณ์โกรธของผู้เรียน และแสดงความรู้สึกของตนเองอย่างซื่อตรงคงเส้นคงวา (Consistent) ทำให้สัมพันธภาพยืนอยู่บนความจริงแท้"
    )

    # Dialogue 4
    add_h2("ตัวอย่างที่ 4: การยอมรับทางบวกอย่างไม่มีเงื่อนไข (ประเด็นความหลากหลายทางเพศ)")
    add_dialogue_box(
        speaker_roles=["ผู้รับฯ", "ผู้ให้ฯ"],
        dialogue_lines=[
            ("ผู้รับการปรึกษา", "ผมไม่แน่ใจว่าจะเล่าดีไหม ผมกลัวว่าอาจารย์จะรังเกียจหรือมองว่าผมผิดปกติที่ผมมีความรักกับคนเพศเดียวกัน"),
            ("ผู้ให้การปรึกษา", "ครูรู้สึกว่าเธอกล้าหาญมากที่เปิดเผยกับครู สิ่งที่ครูคำนึงคือการช่วยเธอจัดการกับความรู้สึกและการตัดสินใจ การยอมรับของครูที่มีต่อตัวเธอไม่ได้ขึ้นกับความชอบส่วนตัวเรื่องเพศสภาพเลย")
        ],
        analysis_text="ผู้ให้คำปรึกษาแยก 'คุณค่าของบุคคล' ออกจากค่านิยมส่วนตัว ไม่ตัดสิน (Non-judgmental) ให้ความสำคัญกับความทุกข์ใจและความกล้าหาญในการเปิดเผยตนเอง ทำให้เกิดความปลอดภัยทางจิตใจขั้นสูงสุด"
    )

    # ------------------ SECTION 8 & 9 ------------------
    add_h1("8. วิวัฒนาการของชื่อทฤษฎีและแนวโน้มในปัจจุบัน")
    add_body("ทฤษฎีของ Carl Rogers มีการพัฒนาชื่อเรียกและขอบเขตตามยุคสมัย ดังนี้:")
    
    t_evo = doc.add_table(rows=4, cols=3)
    t_evo.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_evo.autofit = False
    t_evo.columns[0].width = Inches(1.8)
    t_evo.columns[1].width = Inches(2.2)
    t_evo.columns[2].width = Inches(2.27)
    
    evo_heads = ["ยุค / ชื่อทฤษฎี", "จุดเน้นและลักษณะสำคัญ", "การนำไปประยุกต์ใช้"]
    for i, h in enumerate(evo_heads):
        c = t_evo.cell(0, i)
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, 100, 100, 120, 120)
        set_cell_borders(c, color="CBD5E1")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_font_run(r, size_pt=14, bold=True, color_rgb=RGBColor(255, 255, 255))
        
    evo_data = [
        ("1. การให้คำปรึกษาแบบไม่นำทาง\n(Non-directive Counseling)", "เน้นการไม่สั่งสอน ไม่ชี้นำ ปล่อยให้ผู้รับบริการเป็นผู้นำทิศทางการพูดคุย", "บำบัดรายบุคคลเบื้องต้น"),
        ("2. การให้คำปรึกษาแบบผู้รับบริการเป็นศูนย์กลาง\n(Client-centered Therapy)", "เน้นความสำคัญของสัมพันธภาพและการสำรวจโลกภายในของผู้รับบริการ", "จิตบำบัดและการปรึกษาเชิงลึก"),
        ("3. การให้คำปรึกษาแบบบุคคลเป็นศูนย์กลาง\n(Person-centered Approach)", "ขยายสู่มุมมองมนุษย์ทุกคน เชื่อมั่นในพลังการเติบโตและการมีปฏิสัมพันธ์", "ขยายสู่การศึกษา การบริหาร และกลุ่มสัมพันธ์")
    ]
    for row_idx, (c0_t, c1_t, c2_t) in enumerate(evo_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for c_idx, text_val in enumerate([c0_t, c1_t, c2_t]):
            cell = t_evo.cell(row_idx, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 120, 120)
            set_cell_borders(cell, color="E2E8F0")
            p = cell.paragraphs[0]
            r = p.add_run(text_val)
            bold_flag = True if c_idx == 0 else False
            set_font_run(r, size_pt=13.5, bold=bold_flag, color_rgb=RGBColor(30, 41, 59))

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(4)
    p_sp.paragraph_format.space_after = Pt(4)

    add_body("ในปัจจุบัน ทฤษฎีนี้ได้รับการผสมผสาน (Integrative Approach) กับแนวคิดอื่นมากขึ้น และผู้ให้คำปรึกษามีบทบาท 'กระฉับกระเฉง' (Active) ในกระบวนการบำบัดมากขึ้น โดยยังคงยึดมั่นใน 3 เงื่อนไขหลัก (Empathy, Genuineness, Unconditional Positive Regard) ไว้อย่างเหนียวแน่น")

    add_h1("9. สรุปประเด็นเน้นย้ำสำหรับเตรียมสอบ (Exam Key Takeaways)")
    add_bullet("ผู้ก่อตั้งคือ Carl Rogers เชื่อมั่นใน 'ศักยภาพมนุษย์' และ 'การยอมรับผิดชอบชีวิตตนเอง'", bold_prefix="1. ผู้ก่อตั้งและจุดยืน: ")
    add_bullet("ความขัดแย้งระหว่างภาพลักษณ์ตน (Self) กับ ประสบการณ์จริง (Experience) นำไปสู่ความเครียด สับสน และอาการตึงเครียดทางร่างกาย", bold_prefix="2. ความไม่สอดคล้อง (Incongruence): ")
    add_bullet("คือการยอมรับตนเองอย่างแท้จริงว่าตนมีทั้งด้านดีและด้านลบ (เป็นทั้งเทวดาและปีศาจ) เมื่อยอมรับได้จะเกิดการพัฒนาพฤติกรรมเชิงบวก", bold_prefix="3. เป้าหมายการปรับภาพลักษณ์: ")
    add_bullet("1) การให้คุณค่าของอินทรีย์, 2) ต้องการการยอมรับจากผู้อื่น, 3) ต้องการยอมรับตนเอง, 4) เงื่อนไขความมีคุณค่า", bold_prefix="4. ลำดับ 4 กระบวนการพัฒนา: ")
    add_bullet("Empathy (เข้าใจลึกซึ้ง), Genuineness/Congruence (จริงแท้ไม่เสแสร้ง), Unconditional Positive Regard (ยอมรับทางบวกไม่มีเงื่อนไข)", bold_prefix="5. หัวใจ 3 ประการ (Core Conditions): ")
    add_bullet("การสะท้อนความรู้สึก (Reflecting) มีพลังและเจาะลึกกว่าการเรียบเรียงถ้อยคำ (Paraphrasing) เพราะทำให้เห็นความคาดหวังในตนเอง", bold_prefix="6. เทคนิคภาษาพูด: ")
    add_bullet("ความเงียบ (Silence) คือช่วงเวลาสำคัญในการไตร่ตรองของผู้รับบริการ ไม่ควรเร่งรีบตัดบท", bold_prefix="7. เทคนิคความเงียบ: ")
    add_bullet("ไม่นำทาง (Non-directive) -> ผู้รับบริการเป็นศูนย์กลาง (Client-centered) -> บุคคลเป็นศูนย์กลาง (Person-centered)", bold_prefix="8. วิวัฒนาการชื่อ: ")

    # Save document
    doc.save(output_path)
    print(f"Successfully saved to: {output_path}")

if __name__ == "__main__":
    out_file = "สรุปบทที่10_การให้คำปรึกษาแบบบุคคลเป็นศูนย์กลาง.docx"
    create_ch10_summary_document(out_file)
